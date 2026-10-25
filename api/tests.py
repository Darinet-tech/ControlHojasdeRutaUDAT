from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.core.management import call_command
from django.test import TestCase
from django.utils import timezone
from rest_framework.test import APIClient

from catalogos.models import Brigada, CausaIncidencia, Centro, Clave, Operario, TipoMantenimiento, TipoOrden, TipoTrabajo
from operacion.tests import BaseOperacion, crear_hoja
from operacion.tests_reportes import EscenarioDosDias


class BaseAPI(TestCase):
    @classmethod
    def setUpTestData(cls):
        call_command("cargar_inicial", verbosity=0)
        User = get_user_model()
        cls.probador = User.objects.create_user("probador", password="clave123")
        cls.probador.groups.add(Group.objects.get(name="Probador"))
        cls.jefe = User.objects.create_user("jefe", password="clave123")
        cls.jefe.groups.add(Group.objects.get(name="Jefe"))
        cls.rep = Brigada.objects.get(nombre="Red de Abonados")
        cls.cable = Brigada.objects.get(nombre="Cable")
        cls.ng = Centro.objects.get(codigo="NG")
        cls.op1 = Operario.objects.create(clave="A1", nombre="Operario Uno", brigada=cls.rep)
        cls.op2 = Operario.objects.create(clave="A2", nombre="Operario Dos", brigada=cls.rep)
        cls.cab1 = Operario.objects.create(clave="C1", nombre="Cablero Uno", brigada=cls.cable)
        cls.op3 = Operario.objects.create(clave="A3", nombre="Operario Tres", brigada=cls.rep)

    def setUp(self):
        self.api = APIClient()
        self.api.force_authenticate(self.probador)

    def clave(self, codigo):
        return Clave.objects.get(codigo=codigo).pk

    def crear_hoja(self, brigada=None, operarios=None, fecha=None):
        r = self.api.post("/api/hojas/", {"fecha": fecha or str(timezone.localdate()), "brigada": (brigada or self.rep).pk,
                                          "centros": [self.ng.pk], "operarios": operarios or [self.op1.pk, self.op2.pk]},
                          format="json")
        self.assertEqual(r.status_code, 201, r.data)
        return r.data["id"]

    def asignar(self, hoja, numero, **extra):
        r = self.api.post(f"/api/hojas/{hoja}/acciones/", {"numero_servicio": numero, **extra}, format="json")
        self.assertEqual(r.status_code, 201, r.data)
        return r.data


class SesionYPermisosTests(BaseAPI):
    def test_sin_sesion_no_se_puede_consultar(self):
        self.assertIn(APIClient().get("/api/hojas/").status_code, (401, 403))

    def test_login_y_datos_del_usuario(self):
        api = APIClient()
        self.assertEqual(api.post("/api/auth/login/", {"username": "probador", "password": "mala"}).status_code, 400)
        r = api.post("/api/auth/login/", {"username": "probador", "password": "clave123"})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.data["grupos"], ["Probador"])
        self.assertFalse(r.data["es_staff"])
        self.assertEqual(api.get("/api/auth/me/").data["username"], "probador")
        self.assertEqual(api.get("/api/auth/me/").data["hoy"], timezone.localdate())  # el servidor dice qué día es hoy
        self.assertEqual(api.post("/api/auth/logout/").status_code, 204)
        self.assertIn(api.get("/api/auth/me/").status_code, (401, 403))

    def test_la_pantalla_web_necesita_el_token_csrf_para_escribir(self):
        api = APIClient(enforce_csrf_checks=True)
        api.post("/api/auth/login/", {"username": "probador", "password": "clave123"})
        token = api.get("/api/auth/me/").cookies["csrftoken"].value
        datos = {"fecha": "2026-10-01", "brigada": self.rep.pk, "centros": [self.ng.pk], "operarios": [self.op1.pk]}
        self.assertEqual(api.post("/api/hojas/", datos, format="json").status_code, 403)  # sin token
        self.assertEqual(api.post("/api/hojas/", datos, format="json", HTTP_X_CSRFTOKEN=token).status_code, 201)

    def test_el_jefe_consulta_pero_no_registra(self):
        self.crear_hoja()
        self.api.force_authenticate(self.jefe)
        self.assertEqual(self.api.get("/api/hojas/").status_code, 200)
        r = self.api.post("/api/hojas/", {"fecha": "2026-10-01", "brigada": self.rep.pk, "centros": [self.ng.pk],
                                          "operarios": [self.op1.pk]}, format="json")
        self.assertEqual(r.status_code, 403)

    def test_solo_quien_gestiona_catalogos_ve_el_acceso_a_la_administracion(self):
        probador_staff = get_user_model().objects.create_user("probador_staff", password="x", is_staff=True)
        probador_staff.groups.add(Group.objects.get(name="Probador"))
        especialista = get_user_model().objects.create_user("especialista", password="x", is_staff=True)
        especialista.groups.add(Group.objects.get(name="Especialista"))
        admin = get_user_model().objects.create_superuser("administrador", "a@a.cu", "x")
        quien = {u.username: u for u in (self.probador, self.jefe, probador_staff, especialista, admin)}
        veredicto = {}
        for nombre, usuario in quien.items():
            self.api.force_authenticate(usuario)
            veredicto[nombre] = self.api.get("/api/auth/me/").data["puede_administrar"]
        self.assertEqual(veredicto, {"probador": False, "jefe": False, "probador_staff": False,
                                     "especialista": True, "administrador": True})

    def test_catalogos_para_la_pantalla_de_captura(self):
        self.assertEqual(len(self.api.get("/api/claves/").data), 177)
        self.assertEqual({c["grupo"] for c in self.api.get("/api/claves/?grupo=Cable").data}, {"Cable"})
        c25r = next(c for c in self.api.get("/api/claves/?q=25-R").data if c["codigo"] == "25-R")
        self.assertEqual(c25r["norma_min"], 45)
        Clave.objects.filter(codigo="25-R").update(activa=False)
        self.assertNotIn("25-R", [c["codigo"] for c in self.api.get("/api/claves/").data])
        self.assertEqual(len(self.api.get("/api/tipos-orden/").data), 12)
        self.assertEqual(len(self.api.get("/api/causas-incidencia/").data), 8)
        self.assertEqual(len(self.api.get(f"/api/operarios/?brigada={self.cable.pk}").data), 1)


class FlujoDelProbadorTests(BaseAPI):
    def test_dia_completo(self):
        hoja = self.crear_hoja()
        a1 = self.asignar(hoja, "46-326407")  # 1) se asigna el servicio a la hoja
        a2 = self.asignar(hoja, "46-111111")
        self.assertEqual((a1["numero"], a2["numero"], a1["estado"]), (1, 2, "EN_CURSO"))

        # 2) el operario concluye: el probador registra las claves y da el OK tras hablar con el cliente
        r = self.api.patch(f"/api/acciones/{a1['id']}/", {"claves": [self.clave("25-R"), self.clave("07-IC")]}, format="json")
        self.assertEqual((r.status_code, r.data["norma_min"], r.data["sugerencia_ok"]), (200, 45, True))
        r = self.api.post(f"/api/acciones/{a1['id']}/ok/", {"cliente_contactado": "Sra. Pérez"}, format="json")
        self.assertEqual((r.status_code, r.data["estado"], r.data["estado_orden"]), (200, "OK", "RESTABLECIDA"))
        self.assertIsNotNone(r.data["hora_fin"])

        # 3) la otra acción no se resuelve y se cierra sin OK
        self.api.patch(f"/api/acciones/{a2['id']}/", {"claves": [self.clave("36")]}, format="json")
        r = self.api.post(f"/api/acciones/{a2['id']}/cerrar-sin-ok/")
        self.assertEqual(r.data["estado"], "SIN_OK")

        # 4) incidencia con contador y cierre del día
        lluvia = CausaIncidencia.objects.get(codigo="100").pk
        r = self.api.post(f"/api/hojas/{hoja}/incidencias/iniciar/", {"causa": lluvia}, format="json")
        self.assertEqual((r.status_code, r.data["activa"]), (201, True))
        self.assertEqual(self.api.get(f"/api/hojas/{hoja}/").data["incidencia_activa"], "100 - Lluvia")
        self.assertEqual(self.api.post(f"/api/hojas/{hoja}/incidencias/finalizar/").data["activa"], False)
        self.assertEqual(self.api.post(f"/api/hojas/{hoja}/incidencias/finalizar/").status_code, 400)  # ya no hay activa

        r = self.api.post(f"/api/hojas/{hoja}/cerrar/")
        self.assertEqual((r.status_code, r.data["no_visitadas"]), (200, []))
        self.assertEqual(r.data["resumen"]["efectivas"], 1)
        self.assertEqual(r.data["resumen"]["efectividad_pct"], 50.0)

        # 5) hoja cerrada: ya no admite cambios
        r = self.api.post(f"/api/hojas/{hoja}/acciones/", {"numero_servicio": "46-999999"}, format="json")
        self.assertEqual(r.status_code, 400)
        self.assertEqual(self.api.patch(f"/api/hojas/{hoja}/", {"observaciones": "x"}, format="json").status_code, 400)
        detalle = self.api.get(f"/api/hojas/{hoja}/").data
        self.assertTrue(detalle["cerrada"])
        self.assertEqual((len(detalle["acciones"]), len(detalle["incidencias"])), (2, 1))

    def test_la_hoja_necesita_operarios(self):
        r = self.api.post("/api/hojas/", {"fecha": "2026-10-01", "brigada": self.rep.pk, "centros": [self.ng.pk],
                                          "operarios": []}, format="json")
        self.assertEqual(r.status_code, 400)

    def test_validaciones_de_la_accion(self):
        hoja = self.crear_hoja()
        a = self.asignar(hoja, "46-222222")
        # cerrar sin claves
        r = self.api.post(f"/api/acciones/{a['id']}/ok/", {"cliente_contactado": "Cliente"}, format="json")
        self.assertEqual(r.status_code, 400)
        self.assertIn("clave de cierre", r.data["detail"][0])
        # el mismo servicio abierto en la misma hoja
        r = self.api.post(f"/api/hojas/{hoja}/acciones/", {"numero_servicio": "46-222222"}, format="json")
        self.assertEqual(r.status_code, 400)
        # OK sin cliente
        self.api.patch(f"/api/acciones/{a['id']}/", {"claves": [self.clave("25-R")]}, format="json")
        self.assertEqual(self.api.post(f"/api/acciones/{a['id']}/ok/", {"cliente_contactado": ""}, format="json").status_code, 400)
        # clave desactivada no se acepta
        Clave.objects.filter(codigo="21-R").update(activa=False)
        r = self.api.patch(f"/api/acciones/{a['id']}/", {"claves": [self.clave("21-R")]}, format="json")
        self.assertEqual(r.status_code, 400)
        # una acción cerrada no se modifica
        self.api.post(f"/api/acciones/{a['id']}/ok/", {"cliente_contactado": "Cliente"}, format="json")
        self.assertEqual(self.api.patch(f"/api/acciones/{a['id']}/", {"observaciones": "x"}, format="json").status_code, 400)


class PendientesYReincidenciasTests(BaseAPI):
    def test_pendiente_pasa_a_cable_y_luego_se_restablece(self):
        hoja_rep = self.crear_hoja()
        a = self.asignar(hoja_rep, "46-326407", tipo_orden=None)
        self.api.patch(f"/api/acciones/{a['id']}/", {"claves": [self.clave("23")]}, format="json")
        r = self.api.post(f"/api/acciones/{a['id']}/pasar-a-brigada/", {"brigada": self.cable.pk}, format="json")
        self.assertEqual((r.data["estado"], r.data["pasada_a_nombre"], r.data["estado_orden"]), ("PASADA_A_BRIGADA", "Cable", "PENDIENTE"))

        pendientes = self.api.get(f"/api/ordenes/pendientes/?brigada={self.cable.pk}").data
        self.assertEqual([o["numero_servicio"] for o in pendientes], ["46-326407"])
        self.assertEqual(self.api.get(f"/api/ordenes/pendientes/?brigada={self.rep.pk}").data, [])

        # la hoja de Cable recibe el mismo servicio: se reutiliza la orden pendiente
        hoja_cable = self.crear_hoja(brigada=self.cable, operarios=[self.cab1.pk])
        b = self.asignar(hoja_cable, "46-326407")
        self.assertEqual(b["orden"], a["orden"])
        # ya está en una hoja abierta: no se vuelve a ofrecer como pendiente
        self.assertEqual(self.api.get(f"/api/ordenes/pendientes/?brigada={self.cable.pk}").data, [])
        self.api.patch(f"/api/acciones/{b['id']}/", {"claves": [self.clave("23-R")]}, format="json")
        r = self.api.post(f"/api/acciones/{b['id']}/ok/", {"cliente_contactado": "Cliente"}, format="json")
        self.assertEqual(r.data["estado_orden"], "RESTABLECIDA")
        self.assertEqual(self.api.get(f"/api/ordenes/pendientes/?brigada={self.cable.pk}").data, [])

    def test_historial_y_reincidencia_con_los_operarios_que_trabajaron_el_servicio(self):
        hoja1 = self.crear_hoja(operarios=[self.op1.pk])
        a = self.asignar(hoja1, "46-555555")
        self.api.patch(f"/api/acciones/{a['id']}/", {"claves": [self.clave("25-R")]}, format="json")
        self.api.post(f"/api/acciones/{a['id']}/ok/", {"cliente_contactado": "Cliente"}, format="json")
        r = self.api.get("/api/ordenes/historial/?numero=46-555555").data
        self.assertEqual((r["ordenes"], r["acciones"], r["reincidencia"]), (1, 1, False))

        # el servicio se vuelve a dañar: como la orden anterior ya estaba restablecida, se crea una nueva
        hoja2 = self.crear_hoja(operarios=[self.op2.pk])
        b = self.asignar(hoja2, "46-555555")
        self.assertNotEqual(b["orden"], a["orden"])
        r = self.api.get("/api/ordenes/historial/?numero=46-555555").data
        self.assertEqual((r["ordenes"], r["acciones"], r["reincidencia"]), (2, 2, True))
        self.assertEqual([h["operarios"] for h in r["historial"]], [["Operario Uno"], ["Operario Dos"]])
        self.assertEqual(self.api.get("/api/ordenes/historial/").status_code, 400)

    def test_filtros_de_hojas(self):
        self.crear_hoja()
        self.crear_hoja(brigada=self.cable, operarios=[self.cab1.pk])
        hoy = str(timezone.localdate())
        self.assertEqual(self.api.get(f"/api/hojas/?fecha={hoy}").data["count"], 2)
        self.assertEqual(self.api.get(f"/api/hojas/?fecha={hoy}&brigada={self.cable.pk}").data["count"], 1)
        self.assertEqual(self.api.get("/api/hojas/?fecha=2020-01-01").data["count"], 0)


class ApiNavegableTests(BaseAPI):
    """Quien ya entró a la administración puede abrir la API en el navegador sin configurar nada más."""

    def test_paginas_html_de_la_api(self):
        self.client.force_login(self.probador)
        self.crear_hoja()
        for url in ("/api/", "/api/hojas/", "/api/claves/", "/api/ordenes/", "/api/acciones/", "/api/auth/me/"):
            r = self.client.get(url, HTTP_ACCEPT="text/html")
            self.assertEqual(r.status_code, 200, url)
        hoja = self.client.get("/api/hojas/", HTTP_ACCEPT="application/json").json()["results"][0]["id"]
        self.assertEqual(self.client.get(f"/api/hojas/{hoja}/", HTTP_ACCEPT="text/html").status_code, 200)


class OperarioUnaHojaPorDiaAPITests(BaseAPI):
    def test_no_se_puede_crear_otra_hoja_con_un_operario_ocupado(self):
        self.crear_hoja(operarios=[self.op1.pk], fecha="2026-10-05")
        r = self.api.post("/api/hojas/", {"fecha": "2026-10-05", "brigada": self.cable.pk, "centros": [self.ng.pk],
                                          "operarios": [self.op1.pk, self.op2.pk]}, format="json")
        self.assertEqual(r.status_code, 400)
        self.assertIn("Operario Uno ya está en la hoja NG · Red de Abonados", r.data["operarios"][0])
        self.assertEqual(self.api.get("/api/hojas/?fecha=2026-10-05").data["count"], 1)  # no quedó nada a medias
        self.crear_hoja(operarios=[self.op1.pk], fecha="2026-10-06")  # otro día: permitido

    def test_tampoco_al_modificar_la_hoja(self):
        h1 = self.crear_hoja(operarios=[self.op1.pk], fecha="2026-10-05")
        h2 = self.crear_hoja(operarios=[self.op2.pk], fecha="2026-10-05")
        r = self.api.patch(f"/api/hojas/{h2}/", {"operarios": [self.op1.pk]}, format="json")
        self.assertEqual(r.status_code, 400)
        h3 = self.crear_hoja(operarios=[self.op1.pk], fecha="2026-10-07")
        r = self.api.patch(f"/api/hojas/{h3}/", {"fecha": "2026-10-05"}, format="json")  # mover la hoja a un día ocupado
        self.assertEqual(r.status_code, 400)
        self.assertEqual(self.api.patch(f"/api/hojas/{h1}/", {"observaciones": "ok"}, format="json").status_code, 200)
        self.assertEqual(self.api.patch(f"/api/hojas/{h2}/", {"operarios": [self.op2.pk, self.op3.pk]}, format="json").status_code, 200)


class NoVisitadasAPITests(BaseAPI):
    def test_cierre_con_servicios_no_visitados_y_traerlos_al_dia_siguiente(self):
        hoja1 = self.crear_hoja(operarios=[self.op1.pk, self.op2.pk], fecha="2026-10-05")
        a1 = self.asignar(hoja1, "46-700001"); a2 = self.asignar(hoja1, "46-700002"); a3 = self.asignar(hoja1, "46-700003")
        self.api.patch(f"/api/acciones/{a1['id']}/", {"claves": [self.clave("25-R")]}, format="json")
        self.api.post(f"/api/acciones/{a1['id']}/ok/", {"cliente_contactado": "Cliente"}, format="json")
        self.api.patch(f"/api/acciones/{a2['id']}/", {"claves": [self.clave("23")]}, format="json")
        self.api.post(f"/api/acciones/{a2['id']}/pasar-a-brigada/", {"brigada": self.cable.pk}, format="json")

        r = self.api.post(f"/api/hojas/{hoja1}/cerrar/")
        self.assertEqual(r.data["no_visitadas"], ["46-700003"])
        res = r.data["resumen"]
        self.assertEqual(res["resultados"], {"ok": 1, "sin_ok": 0, "pasadas": 1, "no_visitadas": 1, "en_curso": 0})
        self.assertEqual(res["pasadas_a_brigada"], [{"servicio": "46-700002", "brigada": "Cable"}])
        detalle = self.api.get(f"/api/hojas/{hoja1}/").data
        self.assertEqual([a["estado"] for a in detalle["acciones"]], ["OK", "PASADA_A_BRIGADA", "NO_VISITADA"])
        self.assertEqual(self.api.get("/api/hojas/?fecha=2026-10-05").data["results"][0]["resultados"]["no_visitadas"], 1)

        # al día siguiente la brigada ve lo que quedó sin visitar y lo trae con un clic
        ofrecidas = self.api.get(f"/api/ordenes/no-visitadas/?brigada={self.rep.pk}&centro={self.ng.pk}").data
        self.assertEqual([(o["numero_servicio"], o["ultima_fecha"], o["ultimo_centro"]) for o in ofrecidas],
                         [("46-700003", "2026-10-05", "NG")])
        self.assertEqual(self.api.get(f"/api/ordenes/no-visitadas/?brigada={self.cable.pk}").data, [])
        self.assertEqual(self.api.get("/api/ordenes/no-visitadas/").status_code, 400)

        hoja2 = self.crear_hoja(operarios=[self.op1.pk, self.op2.pk], fecha="2026-10-06")
        r = self.api.post(f"/api/hojas/{hoja2}/asignar-no-visitadas/")
        self.assertEqual((r.status_code, [a["numero_servicio"] for a in r.data]), (201, ["46-700003"]))
        self.assertEqual(r.data[0]["orden"], a3["orden"])
        self.assertEqual(r.data[0]["estado_orden"], "ABIERTA")
        self.assertEqual(self.api.get(f"/api/ordenes/no-visitadas/?brigada={self.rep.pk}&centro={self.ng.pk}").data, [])
        self.assertEqual(self.api.post(f"/api/hojas/{hoja2}/asignar-no-visitadas/").data, [])

    def test_escribir_el_numero_de_un_servicio_no_visitado_reutiliza_su_orden(self):
        hoja1 = self.crear_hoja(operarios=[self.op1.pk], fecha="2026-10-05")
        a = self.asignar(hoja1, "46-800001")
        self.api.post(f"/api/hojas/{hoja1}/cerrar/")
        hoja2 = self.crear_hoja(operarios=[self.op1.pk], fecha="2026-10-06")
        b = self.asignar(hoja2, "46-800001")
        self.assertEqual(b["orden"], a["orden"])
        self.assertEqual(b["estado_orden"], "ABIERTA")
        self.assertEqual(self.api.get(f"/api/ordenes/no-visitadas/?brigada={self.rep.pk}").data, [])


class ResumenAPITests(EscenarioDosDias, BaseOperacion):
    def setUp(self):
        self.escenario()
        self.api = APIClient()
        self.jefe = get_user_model().objects.create_user("jefe_resumen", password="x")
        self.jefe.groups.add(Group.objects.get(name="Jefe"))

    def test_hace_falta_iniciar_sesion(self):
        self.assertIn(self.api.get("/api/resumen/?periodo=mes&fecha=2026-10-01").status_code, (401, 403))

    def test_el_jefe_y_el_probador_consultan_el_resumen(self):
        for usuario in (self.jefe, self.user):
            if usuario is self.user:
                usuario.groups.add(Group.objects.get(name="Probador"))
            self.api.force_authenticate(usuario)
            r = self.api.get("/api/resumen/?periodo=mes&fecha=2026-10-15")
            self.assertEqual(r.status_code, 200)
            self.assertEqual((r.data["totales"]["hojas"], r.data["totales"]["efectividad_pct"], r.data["totales"]["productividad_pct"]),
                             (2, 50.0, 33.3))

    def test_dia_con_filtros_y_listados(self):
        self.api.force_authenticate(self.jefe)
        r = self.api.get(f"/api/resumen/?periodo=dia&fecha=2026-10-05&brigada={self.rep.pk}&centro={self.centro.pk}")
        self.assertEqual([h["id"] for h in r.data["hojas"]], [self.hoja_a.pk])
        self.assertEqual(r.data["pasadas_a_brigada"]["lista"][0]["hacia"], "Cable")
        self.assertEqual(r.data["no_visitadas"]["total"], 1)
        self.assertEqual(self.api.get(f"/api/resumen/?periodo=dia&fecha=2026-10-05&brigada={self.cable.pk}").data["totales"]["hojas"], 0)

    def test_sin_parametros_usa_el_dia_de_hoy_y_valida_los_datos(self):
        self.api.force_authenticate(self.jefe)
        r = self.api.get("/api/resumen/")
        self.assertEqual((r.status_code, r.data["periodo"], r.data["desde"]), (200, "dia", str(timezone.localdate())))
        self.assertEqual(self.api.get("/api/resumen/?periodo=semana").status_code, 400)
        self.assertEqual(self.api.get("/api/resumen/?fecha=05-10-2026").status_code, 400)


class VariosCentrosAPITests(BaseAPI):
    def setUp(self):
        super().setUp()
        self.fi = Centro.objects.get(codigo="FI")
        self.fo = Centro.objects.get(codigo="FO")

    def hoja_con(self, centros, operarios=None, fecha="2026-10-05"):
        r = self.api.post("/api/hojas/", {"fecha": fecha, "brigada": self.rep.pk, "centros": [c.pk for c in centros],
                                          "operarios": operarios or [self.op1.pk]}, format="json")
        self.assertEqual(r.status_code, 201, r.data)
        return r.data

    def test_una_hoja_puede_tener_varios_centros(self):
        hoja = self.hoja_con([self.ng, self.fi])
        self.assertEqual((hoja["centros_codigos"], hoja["centros_etiqueta"]), (["FI", "NG"], "FI + NG"))
        for centro in (self.ng, self.fi):  # el filtro encuentra la hoja por cualquiera de sus centros
            self.assertEqual(self.api.get(f"/api/hojas/?fecha=2026-10-05&centro={centro.pk}").data["count"], 1)
        self.assertEqual(self.api.get(f"/api/hojas/?fecha=2026-10-05&centro={self.fo.pk}").data["count"], 0)

    def test_la_hoja_necesita_al_menos_un_centro(self):
        r = self.api.post("/api/hojas/", {"fecha": "2026-10-05", "brigada": self.rep.pk, "centros": [],
                                          "operarios": [self.op1.pk]}, format="json")
        self.assertEqual(r.status_code, 400)
        self.assertIn("centro", r.data["centros"][0].lower())

    def test_el_servicio_indica_su_centro_cuando_la_hoja_tiene_varios(self):
        hoja = self.hoja_con([self.ng, self.fi])["id"]
        r = self.api.post(f"/api/hojas/{hoja}/acciones/", {"numero_servicio": "46-900001"}, format="json")
        self.assertEqual(r.status_code, 400)
        self.assertIn("indique el centro del servicio", r.data["detail"][0])
        r = self.api.post(f"/api/hojas/{hoja}/acciones/", {"numero_servicio": "46-900001", "centro": self.fo.pk}, format="json")
        self.assertEqual(r.status_code, 400)  # La Reforma no es uno de los centros de la hoja
        r = self.api.post(f"/api/hojas/{hoja}/acciones/", {"numero_servicio": "46-900001", "centro": self.fi.pk}, format="json")
        self.assertEqual((r.status_code, r.data["centro"], r.data["centro_codigo"]), (201, self.fi.pk, "FI"))

    def test_en_una_hoja_de_un_centro_el_servicio_toma_ese_centro(self):
        hoja = self.hoja_con([self.ng])["id"]
        self.assertEqual(self.asignar(hoja, "46-900002")["centro_codigo"], "NG")

    def test_los_no_visitados_se_ofrecen_por_cualquiera_de_los_centros_de_la_hoja(self):
        h1 = self.hoja_con([self.fi], fecha="2026-10-05")["id"]
        self.asignar(h1, "46-900003")
        self.api.post(f"/api/hojas/{h1}/cerrar/")
        base = f"/api/ordenes/no-visitadas/?brigada={self.rep.pk}"
        self.assertEqual(len(self.api.get(f"{base}&centro={self.fi.pk}").data), 1)
        self.assertEqual(len(self.api.get(f"{base}&centro={self.ng.pk}").data), 0)
        self.assertEqual(len(self.api.get(f"{base}&centro={self.ng.pk}&centro={self.fi.pk}").data), 1)  # varios centros
        self.assertEqual(self.api.get(f"{base}&centro={self.fi.pk}").data[0]["ultimo_centro"], "FI")
        h2 = self.hoja_con([self.ng, self.fi], operarios=[self.op2.pk], fecha="2026-10-06")["id"]
        r = self.api.post(f"/api/hojas/{h2}/asignar-no-visitadas/")
        self.assertEqual([(a["numero_servicio"], a["centro_codigo"]) for a in r.data], [("46-900003", "FI")])


class DescargaDeInformesAPITests(EscenarioDosDias, BaseOperacion):
    def setUp(self):
        self.escenario()
        self.api = APIClient()
        self.jefe = get_user_model().objects.create_user("jefe_informe", password="x")
        self.jefe.groups.add(Group.objects.get(name="Jefe"))

    def test_hace_falta_iniciar_sesion(self):
        self.assertIn(self.api.get("/api/informe-diario/?fecha=2026-10-05").status_code, (401, 403))
        self.assertIn(self.api.get(f"/api/hojas/{self.hoja_a.pk}/informe/").status_code, (401, 403))

    def test_el_pdf_se_abre_en_el_navegador_y_el_excel_se_descarga(self):
        self.api.force_authenticate(self.jefe)
        pdf = self.api.get("/api/informe-diario/?fecha=2026-10-05&formato=pdf")
        self.assertEqual((pdf.status_code, pdf["Content-Type"]), (200, "application/pdf"))
        self.assertEqual(pdf["Content-Disposition"], 'inline; filename="UDAT_informe_diario_2026-10-05.pdf"')
        self.assertTrue(pdf.content.startswith(b"%PDF"))
        xlsx = self.api.get("/api/informe-diario/?fecha=2026-10-05&formato=xlsx")
        self.assertEqual(xlsx["Content-Disposition"], 'attachment; filename="UDAT_informe_diario_2026-10-05.xlsx"')
        self.assertTrue(xlsx.content.startswith(b"PK"))  # un .xlsx es un zip
        self.assertEqual(self.api.get("/api/informe-diario/?fecha=2026-10-05").status_code, 200)  # por defecto, PDF

    def test_informe_de_una_hoja(self):
        import io

        import openpyxl
        self.api.force_authenticate(self.jefe)
        r = self.api.get(f"/api/hojas/{self.hoja_a.pk}/informe/?formato=xlsx")
        self.assertEqual(r.status_code, 200)
        self.assertIn(f"_{self.hoja_a.pk}.xlsx", r["Content-Disposition"])
        ws = openpyxl.load_workbook(io.BytesIO(r.content))["Hojas de ruta"]
        self.assertEqual((ws.cell(5, 2).value, ws.cell(1, 7).value), ("46 100001 07-IC 25-R", None))
        self.assertEqual(self.api.get(f"/api/hojas/{self.hoja_a.pk}/informe/").status_code, 200)

    def test_filtros_y_validaciones(self):
        import io

        import openpyxl
        self.api.force_authenticate(self.jefe)
        r = self.api.get(f"/api/informe-diario/?fecha=2026-10-05&formato=xlsx&brigada={self.cable.pk}")
        self.assertEqual(openpyxl.load_workbook(io.BytesIO(r.content))["Resumen"]["A7"].value, 0)  # ninguna hoja de Cable
        self.assertEqual(self.api.get("/api/informe-diario/?fecha=2026-10-05&formato=docx").status_code, 400)
        self.assertEqual(self.api.get("/api/informe-diario/?fecha=05-10-2026").status_code, 400)
        self.assertEqual(self.api.get("/api/hojas/999999/informe/").status_code, 404)


class InicioAPITests(BaseOperacion):
    def setUp(self):
        from datetime import timedelta

        self.api = APIClient()
        self.hoy = timezone.localdate()
        self.ayer = self.hoy - timedelta(days=1)
        self.api.force_authenticate(self.user)
        self.user.groups.add(Group.objects.get(name="Probador"))

    def test_lo_que_requiere_atencion(self):
        from operacion.services import cerrar_hoja, iniciar_incidencia, pasar_a_brigada
        # ayer: una hoja que nadie cerró (con un servicio en curso) y otra cerrada con un servicio no visitado
        sin_cerrar = crear_hoja(fecha=self.ayer, brigada=self.rep, centro=self.centro)
        sin_cerrar.operarios.set([self.op1])
        self.accion(sin_cerrar, "46 10", [], numero=1)
        cerrada = crear_hoja(fecha=self.ayer, brigada=self.cable, centro=self.centro)
        cerrada.operarios.set([self.op2])
        self.accion(cerrada, "46 11", [], numero=1)
        cerrar_hoja(cerrada, self.user)
        # hoy: una hoja con incidencia activa y un servicio pasado a Cable
        hoja = crear_hoja(fecha=self.hoy, brigada=self.rep, centro=self.centro)
        hoja.operarios.set([self.op3])
        a = self.accion(hoja, "46 12", ["23"], numero=1)
        pasar_a_brigada(a, self.cable, self.user)
        iniciar_incidencia(hoja, CausaIncidencia.objects.get(codigo="106"), self.user)

        r = self.api.get("/api/inicio/").data
        self.assertEqual(r["fecha"], self.hoy)
        self.assertEqual((r["hoy"]["hojas"], r["hoy"]["acciones"], r["hoy"]["pasadas"]), (1, 1, 1))
        self.assertEqual([(h["id"], h["en_curso"]) for h in r["hojas_abiertas_anteriores"]], [(sin_cerrar.pk, 1)])
        self.assertEqual([(i["causa"], i["hoja"]) for i in r["incidencias_activas"]], [("106 - Problemas con el transporte", hoja.pk)])
        self.assertEqual([(n["brigada"], n["cantidad"]) for n in r["no_visitadas"]], [("Cable", 1)])
        self.assertEqual([(n["brigada"], n["cantidad"]) for n in r["pendientes"]], [("Cable", 1)])
        self.assertEqual(r["mes"]["hojas"] >= 1, True)
        self.assertEqual([h["id"] for h in r["hojas_hoy"]], [hoja.pk])

    def test_dia_sin_nada_pendiente(self):
        r = self.api.get("/api/inicio/").data
        self.assertEqual((r["hoy"]["hojas"], r["hojas_abiertas_anteriores"], r["incidencias_activas"], r["no_visitadas"], r["pendientes"]), (0, [], [], [], []))

    def test_el_administrador_ve_ademas_el_estado_de_los_catalogos(self):
        from catalogos.models import GrupoClave
        self.assertIsNone(self.api.get("/api/inicio/").data["catalogos"])  # el probador no
        admin = get_user_model().objects.create_superuser("adm_inicio", "a@a.cu", "x")
        self.api.force_authenticate(admin)
        self.assertEqual(self.api.get("/api/inicio/").data["catalogos"], {
            "claves_por_revisar": 5, "claves_sin_norma": 0, "trabajos_sin_norma": 2,  # los trabajos de las órdenes de trabajo aún no tienen norma
            "informe_diario": {"destinatarios": 0, "ultimo": None}})
        Clave.objects.create(codigo="99-Z", grupo=GrupoClave.objects.first(), clasificacion="EFECTIVA")  # una clave nueva sin norma
        self.assertEqual(self.api.get("/api/inicio/").data["catalogos"]["claves_sin_norma"], 1)

    def test_hace_falta_iniciar_sesion(self):
        self.assertIn(APIClient().get("/api/inicio/").status_code, (401, 403))


class CategoriasAPITests(BaseAPI):
    def setUp(self):
        super().setUp()
        self.mppe = TipoMantenimiento.objects.get(codigo="MPPE")
        self.ta = TipoOrden.objects.get(codigo="TA")
        self.poste = TipoTrabajo.objects.get(nombre="Sustitución de poste(s)")
        self.caja = TipoTrabajo.objects.get(nombre="Sustitución de caja(s) terminal(es)")
        self.hoja = self.crear_hoja()

    def post(self, **datos):
        return self.api.post(f"/api/hojas/{self.hoja}/acciones/", datos, format="json")

    def test_cada_tipo_de_accion_pide_sus_datos(self):
        r = self.post(categoria="CORRECTIVO", numero_servicio="46-1")
        self.assertEqual((r.status_code, r.data["etiqueta"], r.data["categoria"], r.data["categoria_texto"]), (201, "46-1", "CORRECTIVO", "Correctivo"))
        r = self.post(categoria="PREVENTIVO", numero_servicio="46-2", tipo_mantenimiento=self.mppe.pk)
        self.assertEqual((r.status_code, r.data["etiqueta"], r.data["tipo_mantenimiento"]), (201, "MPPE 46-2", "MPPE"))
        r = self.post(categoria="ORDEN_SERVICIO", tipo_orden=self.ta.pk, numero_orden="7001", numero_servicio="46-3")
        self.assertEqual((r.status_code, r.data["etiqueta"], r.data["categoria_texto"], r.data["tipo_orden"]), (201, "TA 7001 46-3", "Orden de servicio", "TA"))
        r = self.post(categoria="ORDEN_TRABAJO", numero_orden="8001", trabajos=[self.poste.pk, self.caja.pk])
        self.assertEqual((r.status_code, r.data["etiqueta"], r.data["numero_servicio"], len(r.data["trabajos"])), (201, "OT 8001", "", 2))
        self.assertEqual(self.post(numero_servicio="46-4").data["categoria"], "CORRECTIVO")  # sin indicar tipo: correctivo

    def test_no_se_aceptan_campos_que_no_corresponden_al_tipo(self):
        casos = [
            (dict(categoria="PREVENTIVO", numero_servicio="46-2"), "tipo de mantenimiento"),
            (dict(categoria="CORRECTIVO", numero_servicio="46-1", numero_orden="5"), "no corresponde indicar: el número de orden"),
            (dict(categoria="ORDEN_SERVICIO", numero_servicio="46-3"), "el tipo de orden"),
            (dict(categoria="ORDEN_TRABAJO", numero_orden="8001", trabajos=[self.poste.pk], numero_servicio="46-9"), "no corresponde indicar: el número del servicio"),
            (dict(categoria="ORDEN_TRABAJO", numero_orden="8001"), "los trabajos"),
        ]
        for datos, mensaje in casos:
            r = self.post(**datos)
            self.assertEqual(r.status_code, 400, datos)
            self.assertIn(mensaje, r.data["detail"][0])
        self.assertEqual(self.api.get(f"/api/hojas/{self.hoja}/").data["cantidad_acciones"], 0)

    def test_la_orden_de_trabajo_se_cierra_con_trabajos_y_sin_cliente(self):
        a = self.post(categoria="ORDEN_TRABAJO", numero_orden="8001", trabajos=[self.poste.pk]).data["id"]
        self.assertEqual(self.api.patch(f"/api/acciones/{a}/", {"claves": [self.clave("25-R")]}, format="json").status_code, 400)  # no lleva claves
        r = self.api.patch(f"/api/acciones/{a}/", {"trabajos": [self.caja.pk]}, format="json")
        self.assertEqual((r.status_code, [t["nombre"] for t in r.data["trabajos"]]), (200, [self.caja.nombre]))
        r = self.api.post(f"/api/acciones/{a}/ok/", {}, format="json")
        self.assertEqual((r.status_code, r.data["estado"], r.data["estado_orden"]), (200, "OK", "RESTABLECIDA"))

    def test_el_correctivo_pide_el_cliente_y_el_preventivo_no(self):
        c = self.post(categoria="CORRECTIVO", numero_servicio="46-1").data["id"]
        self.assertEqual(self.api.patch(f"/api/acciones/{c}/", {"trabajos": [self.poste.pk]}, format="json").status_code, 400)  # solo en órdenes de trabajo
        self.api.patch(f"/api/acciones/{c}/", {"claves": [self.clave("25-R")]}, format="json")
        self.assertEqual(self.api.post(f"/api/acciones/{c}/ok/", {}, format="json").status_code, 400)
        p = self.post(categoria="PREVENTIVO", numero_servicio="46-2", tipo_mantenimiento=self.mppe.pk).data["id"]
        self.api.patch(f"/api/acciones/{p}/", {"claves": [self.clave("21-R")]}, format="json")
        self.assertEqual(self.api.post(f"/api/acciones/{p}/ok/", {}, format="json").status_code, 200)

    def test_catalogos_para_el_formulario(self):
        self.assertEqual([t["codigo"] for t in self.api.get("/api/tipos-mantenimiento/").data], ["MPI", "MPPE", "MPPI"])
        trabajos = self.api.get("/api/tipos-trabajo/").data
        self.assertEqual([(t["nombre"], t["norma_min"]) for t in trabajos],
                         [("Sustitución de caja(s) terminal(es)", None), ("Sustitución de poste(s)", None)])  # los especialistas definen las normas

    def test_traer_los_pendientes_a_la_hoja_de_la_brigada(self):
        fi = Centro.objects.get(codigo="FI")
        for numero, centro in (("46-6201", self.ng), ("46-6202", fi)):
            hoja = self.api.post("/api/hojas/", {"fecha": "2026-10-05", "brigada": self.rep.pk, "centros": [centro.pk],
                                                 "operarios": [self.op1.pk if centro == self.ng else self.op2.pk]}, format="json").data["id"]
            a = self.api.post(f"/api/hojas/{hoja}/acciones/", {"numero_servicio": numero}, format="json").data["id"]
            self.api.patch(f"/api/acciones/{a}/", {"claves": [self.clave("23")]}, format="json")
            self.api.post(f"/api/acciones/{a}/pasar-a-brigada/", {"brigada": self.cable.pk}, format="json")
        base = f"/api/ordenes/pendientes/?brigada={self.cable.pk}"
        self.assertEqual(len(self.api.get(base).data), 2)
        self.assertEqual([o["etiqueta"] for o in self.api.get(f"{base}&centro={self.ng.pk}").data], ["46-6201"])
        hoja_cable = self.api.post("/api/hojas/", {"fecha": "2026-10-06", "brigada": self.cable.pk, "centros": [self.ng.pk],
                                                   "operarios": [self.cab1.pk]}, format="json").data["id"]
        r = self.api.post(f"/api/hojas/{hoja_cable}/asignar-pendientes/")
        self.assertEqual((r.status_code, [a["etiqueta"] for a in r.data], [a["centro_codigo"] for a in r.data]), (201, ["46-6201"], ["NG"]))
        self.assertEqual(self.api.post(f"/api/hojas/{hoja_cable}/asignar-pendientes/").data, [])


class AnalisisAPITests(EscenarioDosDias, BaseOperacion):
    def setUp(self):
        from catalogos.models import Vehiculo
        self.escenario()
        self.hoja_a.vehiculo = Vehiculo.objects.create(chapa="B000001")
        self.hoja_a.save()
        self.api = APIClient()
        jefe = get_user_model().objects.create_user("jefe_analisis", password="x")
        jefe.groups.add(Group.objects.get(name="Jefe"))
        self.api.force_authenticate(jefe)

    def consulta(self, periodo):
        from datetime import date
        from unittest import mock
        with mock.patch("django.utils.timezone.localdate", return_value=date(2026, 10, 20)):
            return self.api.get(f"/api/analisis/?periodo={periodo}")

    def test_jornada_anterior_es_el_ultimo_dia_con_hojas(self):
        r = self.consulta("anterior").data
        self.assertEqual((r["fecha"], r["hojas"], r["incidencias"]), ("2026-10-06", 1, []))  # el día 6 no tuvo incidencias
        self.assertEqual([(o["nombre"], o["productividad_pct"]) for o in r["operarios"]], [("Operario Uno", 58.3)])
        self.assertEqual(r["vehiculos"], [])  # la hoja del día 6 no tiene vehículo

    def test_mes_y_anio_con_el_tiempo_perdido_por_causa_y_los_mas_productivos(self):
        for periodo in ("mes", "anio"):
            r = self.consulta(periodo).data
            self.assertEqual(r["hojas"], 2)
            self.assertEqual(r["incidencias"], [{"codigo": "100", "nombre": "Lluvia", "minutos": 60, "hojas": 1}])
            self.assertEqual([(o["nombre"], o["productividad_pct"]) for o in r["operarios"]], [("Operario Uno", 33.3), ("Operario Dos", 25.0)])
            self.assertEqual([(v["nombre"], v["productividad_pct"], v["hojas"]) for v in r["vehiculos"]], [("B000001", 25.0, 1)])

    def test_sin_jornada_anterior_y_validaciones(self):
        from datetime import date
        from unittest import mock
        with mock.patch("django.utils.timezone.localdate", return_value=date(2026, 10, 1)):  # nada antes del 1/10
            r = self.api.get("/api/analisis/?periodo=anterior").data
        self.assertEqual((r["fecha"], r["hojas"], r["operarios"]), (None, 0, []))
        self.assertEqual(self.api.get("/api/analisis/?periodo=semana").status_code, 400)
        self.assertIn(APIClient().get("/api/analisis/").status_code, (401, 403))


class AvisosDelEnvioEnInicioTests(BaseOperacion):
    def setUp(self):
        self.api = APIClient()
        self.admin = get_user_model().objects.create_superuser("adm_envio", "a@a.cu", "x")
        self.api.force_authenticate(self.admin)

    def informe(self):
        return self.api.get("/api/inicio/").data["catalogos"]["informe_diario"]

    def test_el_administrador_ve_los_destinatarios_y_el_ultimo_envio(self):
        from datetime import date
        from operacion.models import DestinatarioInforme, InformeEnviado
        self.assertEqual(self.informe(), {"destinatarios": 0, "ultimo": None})
        DestinatarioInforme.objects.create(nombre="Jefe", correo="jefe@etecsa.cu")
        DestinatarioInforme.objects.create(nombre="Baja", correo="baja@etecsa.cu", activo=False)
        InformeEnviado.objects.create(fecha=date(2026, 10, 5), corte=timezone.now(), estado="ERROR", detalle="gaierror: no resuelve smtp.etecsa.cu")
        InformeEnviado.objects.create(fecha=date(2026, 10, 6), corte=timezone.now(), estado="OMITIDO", detalle="Sin hojas")  # los omitidos no cuentan
        r = self.informe()
        self.assertEqual((r["destinatarios"], r["ultimo"]["estado"], r["ultimo"]["detalle"]), (1, "ERROR", "gaierror: no resuelve smtp.etecsa.cu"))

    def test_el_especialista_ve_las_normas_pero_no_el_envio(self):
        especialista = get_user_model().objects.create_user("esp_envio", password="x", is_staff=True)
        especialista.groups.add(Group.objects.get(name="Especialista"))
        self.api.force_authenticate(especialista)
        catalogos = self.api.get("/api/inicio/").data["catalogos"]
        self.assertEqual((catalogos["trabajos_sin_norma"], catalogos["informe_diario"]), (2, None))
