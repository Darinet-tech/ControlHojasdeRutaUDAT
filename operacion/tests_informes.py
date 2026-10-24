import io
from datetime import date, datetime

from django.utils import timezone

from catalogos.models import Brigada, Clave, Centro
from .informes import generar_informe, hora_corta, reunir_datos
from .models import HojaDeRuta, OrdenServicio
from .services import asignar_orden, cerrar_hoja
from .tests import BaseOperacion, crear_hoja
from .tests_reportes import EscenarioDosDias

try:
    import pdfplumber
except ImportError:  # solo se usa para leer el PDF en las pruebas
    pdfplumber = None

DIA5 = date(2026, 10, 5)


def paginas_del_pdf(contenido):
    with pdfplumber.open(io.BytesIO(contenido)) as pdf:
        return [p.extract_text() or "" for p in pdf.pages]


class FormatoDeHoraTests(BaseOperacion):
    def test_como_en_la_hoja_original(self):
        def en(h, m):
            return timezone.make_aware(datetime(2026, 10, 5, h, m))
        self.assertEqual([hora_corta(en(8, 55)), hora_corta(en(12, 5)), hora_corta(en(15, 0)), hora_corta(en(0, 10)), hora_corta(None)],
                         ["8:55a", "12:05p", "3:00p", "12:10a", ""])


class InformeExcelTests(EscenarioDosDias, BaseOperacion):
    def setUp(self):
        self.escenario()

    def libro(self, **kw):
        import openpyxl
        contenido, tipo, nombre = generar_informe(DIA5, "xlsx", **kw)
        return openpyxl.load_workbook(io.BytesIO(contenido)), tipo, nombre

    def test_tiene_las_tres_hojas_y_el_nombre_del_archivo(self):
        wb, tipo, nombre = self.libro()
        self.assertEqual(wb.sheetnames, ["Resumen", "Hojas de ruta", "Auditoría"])
        self.assertEqual((tipo, nombre), ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "UDAT_informe_diario_2026-10-05.xlsx"))

    def test_las_hojas_de_ruta_siguen_la_disposicion_del_original(self):
        ws = self.libro()[0]["Hojas de ruta"]
        self.assertEqual([ws.cell(1, c).value for c in range(1, 6)], [None, "HOJA DE RUTA", None, None, None])
        # "Centro: x-x" y "Operarios: x-x" en lugar de las etiquetas "Centro" y "pareja" de la hoja manuscrita
        self.assertEqual((ws.cell(2, 1).value, ws.cell(2, 4).value, ws.cell(2, 5).value), ("Centro: NG", "fecha", datetime(2026, 10, 5)))
        self.assertEqual((ws.cell(3, 1).value, ws.cell(3, 4).value), ("Operarios: A1-A2", "chapa"))
        self.assertEqual([ws.cell(4, c).value for c in (1, 2, 4, 5)], ["Trabajo", "Descripción del Trabajo", "Hora", "Cliente"])
        self.assertEqual((ws.cell(5, 1).value, ws.cell(5, 2).value, ws.cell(5, 4).value, ws.cell(5, 5).value), (1, "46 100001 07-IC 25-R", "9:00a", "Cliente"))
        self.assertEqual(ws.cell(6, 2).value, "46 100002 36")  # sin OK: sin cliente
        self.assertIsNone(ws.cell(6, 5).value)
        self.assertEqual(ws.cell(8, 2).value, "46 100004")  # no visitada: sin hora
        self.assertIsNone(ws.cell(8, 4).value)
        self.assertEqual(ws.cell(8, 1).value, 4)
        self.assertIsNone(ws.cell(9, 1).value)  # solo las filas de las acciones del día: no se rellena hasta 30
        self.assertIn("B5:C5", {str(m) for m in ws.merged_cells.ranges})  # la descripción ocupa dos columnas, como en el original

    def test_una_hoja_cada_cinco_columnas(self):
        ws = self.libro()[0]["Hojas de ruta"]
        self.assertEqual((ws.cell(1, 7).value, ws.cell(1, 12).value), (None, None))  # solo hay una hoja el día 5
        wb, *_ = self.libro_dos_hojas()
        ws2 = wb["Hojas de ruta"]
        self.assertEqual((ws2.cell(1, 2).value, ws2.cell(1, 7).value), ("HOJA DE RUTA", "HOJA DE RUTA"))  # segundo bloque desde la columna F
        self.assertEqual(ws2.cell(2, 6).value, "Centro: NG")
        # los bloques salen por brigada: primero Cable (A3) y después Red de Abonados (A1-A2)
        self.assertEqual((ws2.cell(3, 1).value, ws2.cell(3, 6).value), ("Operarios: A3", "Operarios: A1-A2"))

    def libro_dos_hojas(self):
        import openpyxl
        otra = crear_hoja(fecha=DIA5, brigada=self.cable, centro=self.centro)
        otra.operarios.set([self.op3])
        contenido, tipo, nombre = generar_informe(DIA5, "xlsx")
        return openpyxl.load_workbook(io.BytesIO(contenido)), tipo, nombre

    def test_mas_de_30_servicios_continua_la_numeracion(self):
        h = crear_hoja(fecha=date(2026, 10, 8), brigada=self.rep, centro=self.centro)
        h.operarios.set([self.op3])
        for n in range(1, 34):
            asignar_orden(h, OrdenServicio.objects.create(numero_servicio=f"46 9{n:05d}"), self.user)
        import openpyxl
        contenido, *_ = generar_informe(date(2026, 10, 8), "xlsx")
        ws = openpyxl.load_workbook(io.BytesIO(contenido))["Hojas de ruta"]
        self.assertEqual((ws.cell(4 + 33, 1).value, ws.cell(4 + 33, 2).value, ws.cell(4 + 34, 1).value), (33, "46 900033", None))

    def test_la_auditoria_trae_una_fila_por_servicio_con_todos_los_datos(self):
        ws = self.libro()[0]["Auditoría"]
        cab = [c.value for c in ws[1]]
        filas = [dict(zip(cab, [c.value for c in fila])) for fila in ws.iter_rows(min_row=2)]
        self.assertEqual(len(filas), 4)
        ok, sin_ok, pasada, no_visitada = filas
        self.assertEqual((ok["Servicio telefónico"], ok["Claves"], ok["Resultado"], ok["Hora de fin"], ok["Tiempo real (min)"], ok["Norma (min)"], ok["Cliente (certificó OK)"]),
                         ("46 100001", "07-IC 25-R", "OK", "9:00a", 60, 45, "Cliente"))
        self.assertEqual((sin_ok["Resultado"], pasada["Resultado"], no_visitada["Resultado"]), ("Sin OK", "Pasada a Cable", "No visitada"))
        self.assertEqual((no_visitada["Hora de fin"], no_visitada["Tiempo real (min)"], pasada["Hoja cerrada"]), (None, None, "Sí"))
        self.assertEqual((ok["Operarios"], ok["Brigada"], ok["Centro del servicio"], ok["Tipo de acción"], ok["Centros de la hoja"]), ("A1-A2", "Red de Abonados", "NG", "Correctivo", "NG"))

    def test_el_resumen_trae_los_indicadores_del_dia(self):
        ws = self.libro()[0]["Resumen"]
        valores = {ws.cell(6, c).value: ws.cell(7, c).value for c in range(1, 14)}
        self.assertEqual((valores["Hojas"], valores["Servicios"], valores["Con OK"], valores["Pasadas a otra brigada"], valores["No visitadas"]), (1, 4, 1, 1, 1))
        self.assertEqual((valores["Efectividad"], valores["Productividad"], valores["Tiempo real (min)"]), ("33.3 %", "25.0 %", 180))

    def test_una_sola_hoja(self):
        otra = crear_hoja(fecha=DIA5, brigada=self.cable, centro=self.centro)
        otra.operarios.set([self.op3])
        wb, _, nombre = self.libro(solo_hoja=self.hoja_a.pk)
        self.assertIn(f"_{self.hoja_a.pk}.xlsx", nombre)
        self.assertEqual(wb["Hojas de ruta"].cell(1, 7).value, None)  # un solo bloque

    def test_formato_desconocido(self):
        with self.assertRaises(ValueError):
            generar_informe(DIA5, "docx")


class InformePdfTests(EscenarioDosDias, BaseOperacion):
    def setUp(self):
        self.escenario()

    def test_es_un_pdf_valido_con_una_pagina_de_resumen_y_una_por_hoja(self):
        if pdfplumber is None:
            self.skipTest("pdfplumber no está instalado (solo hace falta para leer el PDF en estas pruebas)")
        otra = crear_hoja(fecha=DIA5, brigada=self.cable, centro=self.centro)
        otra.operarios.set([self.op3])
        contenido, tipo, nombre = generar_informe(DIA5, "pdf")
        self.assertTrue(contenido.startswith(b"%PDF"))
        self.assertEqual((tipo, nombre), ("application/pdf", "UDAT_informe_diario_2026-10-05.pdf"))
        paginas = paginas_del_pdf(contenido)
        self.assertEqual(len(paginas), 3)  # resumen + 2 hojas
        for n, texto in enumerate(paginas, start=1):
            self.assertIn("Unidad de Asistencia Técnica", texto)
            self.assertIn(f"Página {n} de 3", texto)

    def test_la_primera_pagina_es_el_resumen_del_dia(self):
        if pdfplumber is None:
            self.skipTest("pdfplumber no está instalado")
        texto = paginas_del_pdf(generar_informe(DIA5, "pdf")[0])[0]
        for esperado in ("Informe diario", "Lunes, 5 de octubre de 2026", "Productividad individual", "33.3 %", "25.0 %",
                         "Pasaron a otra brigada (1)", "46 100003", "No visitadas (1)", "46 100004", "Incidencias por causa", "Lluvia"):
            self.assertIn(esperado, texto)

    def test_cada_hoja_sigue_el_formato_original(self):
        if pdfplumber is None:
            self.skipTest("pdfplumber no está instalado")
        texto = paginas_del_pdf(generar_informe(DIA5, "pdf")[0])[1]
        for esperado in ("HOJA DE RUTA", "Centro: NG", "Operarios: A1-A2", "Fecha: 05/10/2026", "Trabajo", "Descripción del Trabajo", "Hora", "Cliente",
                         "46 100001 07-IC 25-R", "9:00a", "Pasada a Cable", "No visitada", "Elaborada por", "Cerrada el 05/10/2026 15:00"):
            self.assertIn(esperado, texto)
        lineas = texto.splitlines()
        self.assertFalse(any(l.startswith("30 ") or l == "30" for l in lineas))  # solo las filas de las acciones del día, no 30

    def test_una_hoja_con_mas_de_30_servicios_ocupa_otra_pagina(self):
        if pdfplumber is None:
            self.skipTest("pdfplumber no está instalado")
        h = crear_hoja(fecha=date(2026, 10, 9), brigada=self.rep, centro=self.centro)
        h.operarios.set([self.op3])
        for n in range(1, 41):
            asignar_orden(h, OrdenServicio.objects.create(numero_servicio=f"46 8{n:05d}"), self.user)
        paginas = paginas_del_pdf(generar_informe(date(2026, 10, 9), "pdf")[0])
        self.assertGreaterEqual(len(paginas), 3)
        self.assertIn("46 800040", " ".join(paginas))

    def test_un_dia_sin_hojas(self):
        if pdfplumber is None:
            self.skipTest("pdfplumber no está instalado")
        paginas = paginas_del_pdf(generar_informe(date(2026, 10, 20), "pdf")[0])
        self.assertEqual(len(paginas), 1)
        self.assertIn("No hay hojas de ruta para este día", paginas[0])

    def test_informe_de_una_sola_hoja(self):
        if pdfplumber is None:
            self.skipTest("pdfplumber no está instalado")
        contenido, _, nombre = generar_informe(DIA5, "pdf", solo_hoja=self.hoja_a.pk)
        self.assertEqual(len(paginas_del_pdf(contenido)), 2)
        self.assertTrue(nombre.startswith("UDAT_hoja_de_ruta_2026-10-05_NG_Red-de-Abonados_"))

    def test_los_datos_incluyen_quien_genero_el_informe(self):
        datos = reunir_datos(DIA5, usuario=self.user)
        self.assertEqual(datos["generado_por"], "probador1")
        self.assertEqual(datos["hojas"][0]["pareja"], "A1 A2")
