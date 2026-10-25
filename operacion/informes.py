"""Informes imprimibles y auditables: la hoja de ruta con el formato de la hoja original, y el resumen del día.

El PDF y el Excel salen de los mismos datos (`reunir_datos`), así que siempre dicen lo mismo. Cada hoja de ruta conserva las
columnas de la hoja manuscrita (Trabajo, Descripción del Trabajo, Hora, Cliente) y agrega, aparte, las de control
(Resultado, tiempo real, norma) para poder auditar. El Excel además trae una hoja "Auditoría" con una fila por servicio.
"""
import io
import re
from datetime import date
from pathlib import Path

from django.conf import settings
from django.db.models import Prefetch
from django.utils import timezone

from catalogos.models import Clave

from .models import Accion
from .reportes import MAXIMO_LISTADO, Normas, hojas_del_periodo, resumen_periodo

LOGO = Path(settings.BASE_DIR) / "frontend" / "public" / "logo-etecsa.png"
LOGO_UDAT = Path(settings.BASE_DIR) / "frontend" / "public" / "logo-udat.png"
AZUL, VERDE = "#050752", "#059E1A"
MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre"]
DIAS = ["lunes", "martes", "miércoles", "jueves", "viernes", "sábado", "domingo"]
RESULTADO = {"OK": "OK", "SIN_OK": "Sin OK", "NO_VISITADA": "No visitada", "EN_CURSO": "En curso"}


# ---------------------------------------------------------------- utilidades de formato

def fecha_larga(f):
    return f"{DIAS[f.weekday()]}, {f.day} de {MESES[f.month - 1]} de {f.year}"


def hora_corta(dt):
    """Como en la hoja original: 8:55a, 12:05p."""
    if not dt:
        return ""
    local = timezone.localtime(dt)
    return f"{local.hour % 12 or 12}:{local.minute:02d}{'a' if local.hour < 12 else 'p'}"


def minutos_a_texto(m):
    if m is None:
        return "—"
    m = round(m)
    return f"{m // 60} h {m % 60:02d} min" if m >= 60 else f"{m} min"


def pct(v):
    return "—" if v is None else f"{v} %"


def nombre_de(usuario):
    return (usuario.get_full_name() or usuario.username) if usuario else ""


def _limpio(texto):
    return re.sub(r"[^A-Za-z0-9._-]+", "-", texto).strip("-")


# ---------------------------------------------------------------- datos

def reunir_datos(fecha, brigada=None, centro=None, solo_hoja=None, usuario=None, ahora=None, corte=None):
    """Todo lo que lleva un informe del día (o de una sola hoja)."""
    ahora = ahora or timezone.now()
    resumen = resumen_periodo("dia", fecha, brigada=brigada, centro=centro, solo_hoja=solo_hoja, ahora=ahora)
    base = hojas_del_periodo(fecha, fecha, brigada, centro, solo_hoja)
    hojas = list(base.select_related("brigada", "vehiculo", "creado_por", "cerrada_por")
                 .prefetch_related("operarios", "centros", "incidencias__causa").order_by("brigada__nombre", "pk"))
    acciones = (Accion.objects.filter(hoja__in=base)
                .select_related("orden__tipo_orden", "orden__tipo_mantenimiento", "centro", "pasada_a", "creado_por", "certificada_por", "hoja")
                .prefetch_related(Prefetch("claves", queryset=Clave.objects.order_by("codigo")), "trabajos").order_by("hoja_id", "numero"))
    normas = Normas()
    indicadores = {h["id"]: h for h in resumen["hojas"]}
    por_hoja = {h.pk: [] for h in hojas}
    for a in acciones:
        claves, orden = list(a.claves.all()), a.orden
        trabajos = sorted(t.nombre for t in a.trabajos.all())
        es_ot = orden.categoria == "ORDEN_TRABAJO"
        if es_ot:  # una orden de trabajo se norma por sus trabajos
            minutos = [normas.minutos_trabajo(t.pk, a.hoja.fecha) for t in a.trabajos.all()]
            descripcion = f"{orden.etiqueta} · " + ", ".join(trabajos) if trabajos else orden.etiqueta
        else:
            minutos = [normas.minutos(c.pk, orden.tipo_orden_id, a.hoja.fecha) for c in claves]
            descripcion = " ".join([orden.etiqueta] + [c.codigo for c in claves])
        estado = a.estado
        por_hoja[a.hoja_id].append({
            "numero": a.numero, "servicio": orden.numero_servicio, "etiqueta": orden.etiqueta, "categoria": orden.get_categoria_display(),
            "tipo_mantenimiento": orden.tipo_mantenimiento.codigo if orden.tipo_mantenimiento_id else "",
            "tipo": orden.tipo_orden.codigo if orden.tipo_orden_id else "", "numero_orden": orden.numero_orden, "trabajos": trabajos,
            "centro": a.centro.codigo if a.centro else "", "claves": [c.codigo for c in claves], "descripcion": descripcion,
            "hora": hora_corta(a.hora_fin), "cliente": a.cliente_contactado, "estado": estado,
            "resultado": f"Pasada a {a.pasada_a.nombre}" if estado == "PASADA_A_BRIGADA" else RESULTADO[estado],
            "tiempo_real": a.tiempo_real_min, "norma": max([m for m in minutos if m is not None], default=None),
            "descripcion_trabajo": a.descripcion_trabajo, "certificada_por": nombre_de(a.certificada_por),
            "creada_por": nombre_de(a.creado_por),
        })
    detalle = []
    for h in hojas:
        operarios = sorted((o.clave, o.nombre) for o in h.operarios.all())
        detalle.append({
            "id": h.pk, "fecha": h.fecha, "brigada": h.brigada.nombre, "centros": h.etiqueta_centros,
            "centros_claves": "-".join(sorted(c.codigo for c in h.centros.all())), "operarios_claves": "-".join(c for c, _ in operarios),
            "pareja": " ".join(c for c, _ in operarios), "operarios": operarios, "vehiculo": h.vehiculo.chapa if h.vehiculo else "",
            "hora_inicio": h.hora_inicio, "observaciones": h.observaciones, "cerrada": h.cerrada,
            "cerrada_en": h.cerrada_en, "cerrada_por": nombre_de(h.cerrada_por), "creada_por": nombre_de(h.creado_por),
            "acciones": por_hoja[h.pk], "indicadores": indicadores.get(h.pk, {}),
            "incidencias": [{"causa": str(i.causa), "inicio": hora_corta(i.inicio), "fin": hora_corta(i.fin) if i.fin else "en curso",
                             "minutos": i.minutos(hasta=ahora)} for i in h.incidencias.all()],
        })
    return {"fecha": fecha, "resumen": resumen, "hojas": detalle, "generado_en": ahora, "generado_por": nombre_de(usuario),
            "una_hoja": bool(solo_hoja), "corte": corte}


def _filas(hoja):
    """Solo las filas con los servicios introducidos en el día (no se rellena hasta 30 como en la hoja manuscrita)."""
    return [(a["numero"], a) for a in sorted(hoja["acciones"], key=lambda a: a["numero"])]


# ---------------------------------------------------------------- PDF

def _pdf(datos):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    azul, verde = colors.HexColor(AZUL), colors.HexColor(VERDE)
    gris, gris_claro = colors.HexColor("#c3cad5"), colors.HexColor("#eff1f5")
    base = ParagraphStyle("base", fontName="Helvetica", fontSize=9, leading=11.5)
    chico = ParagraphStyle("chico", parent=base, fontSize=8, leading=10)
    celda = ParagraphStyle("celda", parent=base, fontSize=8.5, leading=10)
    titulo = ParagraphStyle("titulo", parent=base, fontName="Helvetica-Bold", fontSize=15, leading=18, textColor=azul)
    h2 = ParagraphStyle("h2", parent=base, fontName="Helvetica-Bold", fontSize=10.5, leading=13, textColor=azul, spaceBefore=9, spaceAfter=3)
    hoja_titulo = ParagraphStyle("hoja", parent=base, fontName="Helvetica-Bold", fontSize=16, leading=20, alignment=TA_CENTER, textColor=azul)
    esc = lambda t: str(t).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")  # noqa: E731
    P = lambda t, st=celda: Paragraph(esc(t), st)  # noqa: E731

    ancho = A4[0] - 30 * mm

    cabecera_celda = ParagraphStyle("cabecera_celda", parent=base, fontName="Helvetica-Bold", fontSize=8.5, leading=10, textColor=colors.white, alignment=TA_CENTER)

    def tabla(filas, anchos, encabezado=True, extra=()):
        if encabezado:  # los títulos de columna pueden ocupar dos líneas sin montarse unos sobre otros
            filas = [[Paragraph(esc(x), cabecera_celda) if isinstance(x, str) else x for x in filas[0]]] + list(filas[1:])
        t = Table(filas, colWidths=anchos, repeatRows=1 if encabezado else 0)
        estilo = [("FONTNAME", (0, 0), (-1, -1), "Helvetica"), ("FONTSIZE", (0, 0), (-1, -1), 8.5), ("GRID", (0, 0), (-1, -1), 0.4, gris),
                  ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 2.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5)]
        if encabezado:
            estilo += [("BACKGROUND", (0, 0), (-1, 0), azul)]
        t.setStyle(TableStyle(estilo + list(extra)))
        return t

    # ---------- primera página: resumen del día
    r, t = datos["resumen"], datos["resumen"]["totales"]
    story = [Paragraph(esc("Informe diario" if not datos["una_hoja"] else "Hoja de ruta"), titulo),
             Paragraph(esc(fecha_larga(datos["fecha"]).capitalize()), base), Spacer(1, 5)]
    kpi = [[P("Hojas de ruta"), P("Servicios"), P("Efectividad"), P("Productividad")],
           [Paragraph(f"<b>{t['hojas']}</b>", titulo), Paragraph(f"<b>{t['acciones']}</b>", titulo),
            Paragraph(f"<b>{esc(pct(t['efectividad_pct']))}</b>", titulo), Paragraph(f"<b>{esc(pct(t['productividad_pct']))}</b>", titulo)],
           [P("Tiempo real"), P("Incidencias"), P("Tiempo disponible"), P("Ocupación")],
           [Paragraph(f"<b>{esc(minutos_a_texto(t['tiempo_real_min']))}</b>", titulo), Paragraph(f"<b>{esc(minutos_a_texto(t['incidencias_min']))}</b>", titulo),
            Paragraph(f"<b>{esc(minutos_a_texto(t['tiempo_disponible_min']))}</b>", titulo), Paragraph(f"<b>{esc(pct(t['ocupacion_pct']))}</b>", titulo)]]
    story.append(tabla(kpi, [ancho / 4] * 4, encabezado=False, extra=[("BACKGROUND", (0, 0), (-1, 0), gris_claro), ("BACKGROUND", (0, 2), (-1, 2), gris_claro),
                                                                       ("TEXTCOLOR", (0, 0), (-1, -1), azul)]))
    story.append(Spacer(1, 4))
    story.append(Paragraph(esc(f"Con OK: {t['ok']}  ·  Sin OK: {t['sin_ok']}  ·  Pasadas a otra brigada: {t['pasadas']}  ·  "
                               f"No visitadas: {t['no_visitadas']}  ·  En curso: {t['en_curso']}"), base))
    story.append(Paragraph(esc("Efectividad = con OK entre atendidas (las no visitadas no cuentan). Productividad = tiempo normado de lo efectivo entre tiempo real."), chico))

    if r["hojas"]:
        story.append(Paragraph("Hojas de ruta del día", h2))
        filas = [["Hoja", "Operarios", "Serv.", "OK / sin OK / pasadas / no visitadas", "Efect.", "Prod."]]
        for h in r["hojas"]:
            filas.append([P(f"{h['centro']} · {h['brigada']}"), P(", ".join(h["operarios"])), h["acciones"],
                          f"{h['ok']} / {h['sin_ok']} / {h['pasadas']} / {h['no_visitadas']}", pct(h["efectividad_pct"]), pct(h["productividad_pct"])])
        story.append(tabla(filas, [44 * mm, 52 * mm, 12 * mm, 40 * mm, 16 * mm, 16 * mm]))

        story.append(Paragraph("Productividad individual", h2))
        story.append(Paragraph(esc("La productividad es de la hoja completa: cada operario recibe la de las hojas en que participó."), chico))
        filas = [["Operario", "Hojas", "Servicios", "Con OK", "Efectividad", "Productividad"]]
        filas += [[P(o["nombre"]), o["hojas"], o["acciones"], o["ok"], pct(o["efectividad_pct"]), pct(o["productividad_pct"])] for o in r["por_operario"]]
        story.append(tabla(filas, [60 * mm, 16 * mm, 20 * mm, 20 * mm, 30 * mm, 34 * mm]))

        if r["por_centro"]:
            story.append(Paragraph("Por centro (según el centro de cada servicio)", h2))
            filas = [["Centro", "Servicios", "Con OK", "Efectividad", "Productividad"]]
            filas += [[c["nombre"], c["acciones"], c["ok"], pct(c["efectividad_pct"]), pct(c["productividad_pct"])] for c in r["por_centro"]]
            story.append(tabla(filas, [40 * mm, 30 * mm, 30 * mm, 40 * mm, 40 * mm]))

        def listado(titulo_, filas_, cab, anchos, vacio):
            story.append(Paragraph(esc(titulo_), h2))
            story.append(tabla([cab] + filas_, anchos) if filas_ else Paragraph(esc(vacio), chico))

        pa = r["pasadas_a_brigada"]
        listado(f"Pasaron a otra brigada ({pa['total']})", [[p["servicio"], P(f"{p['centro']} · {p['desde']}"), P(p["hacia"]), P(p["estado_actual"])] for p in pa["lista"]],
                ["Servicio", "De", "Hacia", "Estado actual"], [34 * mm, 55 * mm, 40 * mm, 51 * mm], "Ningún servicio pasó a otra brigada.")
        nv = r["no_visitadas"]
        listado(f"No visitadas ({nv['total']}) — pasan al siguiente día", [[n["servicio"], P(f"{n['centro']} · {n['brigada']}"), "Sigue sin visitar" if n["sigue_sin_visitar"] else "Ya atendida"] for n in nv["lista"]],
                ["Servicio", "Brigada", "Situación"], [34 * mm, 90 * mm, 56 * mm], "Todos los servicios se visitaron.")
        rn = r["reincidencias"]
        if rn["total"]:
            listado(f"Reincidencias ({rn['total']})", [[x["servicio"], x["ordenes"], P(", ".join(x["operarios"]))] for x in rn["lista"]],
                    ["Servicio", "Órdenes", "Operarios que lo trabajaron"], [34 * mm, 20 * mm, 126 * mm], "")
        if r["incidencias"]:
            listado("Incidencias por causa", [[P(f"{i['codigo']} · {i['nombre']}"), minutos_a_texto(i["minutos"]), i["hojas"]] for i in r["incidencias"]],
                    ["Causa", "Tiempo", "Hojas"], [100 * mm, 40 * mm, 40 * mm], "")
    else:
        story.append(Spacer(1, 12))
        story.append(Paragraph(esc("No hay hojas de ruta para este día."), base))

    # ---------- una página por hoja, con el formato de la hoja original
    for h in datos["hojas"]:
        story.append(PageBreak())
        caja = Table([[Paragraph("HOJA DE RUTA", hoja_titulo)]], colWidths=[ancho])
        caja.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 1, azul), ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
        story.append(caja)
        cab = [[Paragraph(f"<b>Centro:</b> {esc(h['centros_claves'])}", base), Paragraph(f"<b>Fecha:</b> {h['fecha']:%d/%m/%Y}", base)],
               [Paragraph(f"<b>Operarios:</b> {esc(h['operarios_claves'])}", base), Paragraph(f"<b>Vehículo (chapa):</b> {esc(h['vehiculo'] or '—')}", base)],
               [P(f"{h['brigada']} · " + ", ".join(n for _, n in h["operarios"]) + (f" · jornada desde las {h['hora_inicio']:%H:%M}" if h["hora_inicio"] else ""), chico), ""]]
        c = tabla(cab, [ancho * 0.6, ancho * 0.4], encabezado=False, extra=[("SPAN", (0, 2), (-1, 2)), ("BACKGROUND", (0, 0), (-1, 1), gris_claro)])
        story.append(c)
        story.append(Spacer(1, 4))

        filas = [["Trabajo", "Descripción del Trabajo", "Hora", "Cliente", "Resultado", "T. real", "Norma"]]
        estilos = []
        for n, a in _filas(h):
            filas.append([str(n), P(a["descripcion"]), a["hora"], P(a["cliente"]), P(a["resultado"]),
                          "" if a["tiempo_real"] is None else str(a["tiempo_real"]), "" if a["norma"] is None else str(a["norma"])])
            color = {"OK": verde, "SIN_OK": colors.HexColor("#b42318"), "PASADA_A_BRIGADA": colors.HexColor("#8a5a00"),
                     "NO_VISITADA": colors.HexColor("#4b5666"), "EN_CURSO": azul}[a["estado"]]
            estilos.append(("TEXTCOLOR", (4, len(filas) - 1), (4, len(filas) - 1), color))
        if len(filas) == 1:
            filas.append(["", P("Sin acciones registradas"), "", "", "", "", ""])
        anchos = [16 * mm, 55 * mm, 15 * mm, 32 * mm, 34 * mm, 14 * mm, 14 * mm]
        estilos += [("BACKGROUND", (4, 0), (-1, 0), verde), ("ALIGN", (0, 0), (0, -1), "CENTER"), ("ALIGN", (2, 0), (2, -1), "CENTER"),
                    ("ALIGN", (5, 0), (6, -1), "CENTER"), ("LINEBEFORE", (4, 0), (4, -1), 1.2, azul), ("FONTSIZE", (0, 1), (-1, -1), 8.5),
                    ("TOPPADDING", (0, 1), (-1, -1), 3), ("BOTTOMPADDING", (0, 1), (-1, -1), 3)]
        story.append(tabla(filas, anchos, extra=estilos))
        story.append(Paragraph(esc("Las columnas Resultado, T. real y Norma (minutos) son de control: no figuran en la hoja manuscrita."), chico))

        ind = h["indicadores"]
        if ind:
            story.append(Paragraph(esc(f"Servicios: {ind['acciones']} · Con OK: {ind['ok']} · Efectividad: {pct(ind['efectividad_pct'])} · Productividad: "
                                       f"{pct(ind['productividad_pct'])} · Tiempo real: {minutos_a_texto(ind['tiempo_real_min'])} · Incidencias: "
                                       f"{minutos_a_texto(ind['incidencias_min'])}"), base))
        if h["incidencias"]:
            story.append(Paragraph("Incidencias", h2))
            story.append(tabla([["Causa", "Desde", "Hasta", "Tiempo"]] + [[P(i["causa"]), i["inicio"], i["fin"], minutos_a_texto(i["minutos"])] for i in h["incidencias"]],
                               [90 * mm, 28 * mm, 28 * mm, 34 * mm]))
        if h["observaciones"]:
            story.append(Paragraph("Observaciones", h2))
            story.append(Paragraph(esc(h["observaciones"]), base))
        cierre = (f"Cerrada el {timezone.localtime(h['cerrada_en']):%d/%m/%Y %H:%M} por {h['cerrada_por']}" if h["cerrada"] and h["cerrada_en"]
                  else (f"Hoja abierta al corte de las {timezone.localtime(datos['corte']):%H:%M}" if datos["corte"] else "Hoja abierta (no se ha cerrado el día)"))
        story.append(Spacer(1, 4))
        story.append(Paragraph(esc(f"Elaborada por: {h['creada_por'] or '—'}  ·  {cierre}"), chico))

    # ---------- encabezado y pie de cada página
    pie_texto = f"Generado el {timezone.localtime(datos['generado_en']):%d/%m/%Y %H:%M}" + (f" por {datos['generado_por']}" if datos["generado_por"] else "")
    subtitulo = "Informe diario" if not datos["una_hoja"] else "Hoja de ruta"
    if datos["corte"]:
        subtitulo += f" · corte a las {timezone.localtime(datos['corte']):%H:%M}"

    def dibujar_cabecera(lienzo, doc):
        w, hgt = A4
        lienzo.saveState()
        lienzo.setFillColor(azul)
        lienzo.rect(0, hgt - 22 * mm, w, 22 * mm, stroke=0, fill=1)
        lienzo.setFillColor(verde)
        lienzo.rect(0, hgt - 22 * mm - 1.2 * mm, w, 1.2 * mm, stroke=0, fill=1)
        if LOGO.exists():
            lienzo.drawImage(str(LOGO), 15 * mm, hgt - 20 * mm, 17 * mm, 17 * mm, mask="auto", preserveAspectRatio=True)
        if LOGO_UDAT.exists():
            lienzo.drawImage(str(LOGO_UDAT), w - 15 * mm - 16 * mm, hgt - 19.5 * mm, 16 * mm, 16 * mm, mask="auto", preserveAspectRatio=True)
        lienzo.setFillColor(colors.white)
        lienzo.setFont("Helvetica-Bold", 12.5)
        lienzo.drawString(36 * mm, hgt - 10.5 * mm, "Unidad de Asistencia Técnica · UDAT")
        lienzo.setFont("Helvetica", 9)
        lienzo.setFillColor(colors.HexColor("#9fe8aa"))
        lienzo.drawString(36 * mm, hgt - 16 * mm, subtitulo)
        lienzo.restoreState()

    class LienzoNumerado(canvas.Canvas):
        def __init__(self, *a, **k):
            super().__init__(*a, **k)
            self._guardadas = []

        def showPage(self):
            self._guardadas.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            total = len(self._guardadas)
            for estado in self._guardadas:
                self.__dict__.update(estado)
                self.setFont("Helvetica", 7.5)
                self.setFillColor(colors.HexColor("#4b5666"))
                self.drawString(15 * mm, 9 * mm, pie_texto)
                self.drawRightString(A4[0] - 15 * mm, 9 * mm, f"Página {self._pageNumber} de {total}")
                super().showPage()
            super().save()

    memoria = io.BytesIO()
    doc = SimpleDocTemplate(memoria, pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm, topMargin=29 * mm, bottomMargin=15 * mm,
                            title=f"{subtitulo} {datos['fecha']:%d/%m/%Y}", author="UDAT")
    doc.build(story, onFirstPage=dibujar_cabecera, onLaterPages=dibujar_cabecera, canvasmaker=LienzoNumerado)
    return memoria.getvalue()


# ---------------------------------------------------------------- Excel

def _excel(datos):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    fino = Side(style="thin", color="999999")
    borde = Border(left=fino, right=fino, top=fino, bottom=fino)
    azul = PatternFill("solid", fgColor=AZUL[1:])
    claro = PatternFill("solid", fgColor="E6E8EE")
    blanco_negrita = Font(bold=True, color="FFFFFF")
    centro = Alignment(horizontal="center", vertical="center", wrap_text=True)
    r, t = datos["resumen"], datos["resumen"]["totales"]

    def encabezado(ws, fila, textos, col=1):
        for i, texto in enumerate(textos):
            c = ws.cell(fila, col + i, texto)
            c.fill, c.font, c.alignment, c.border = azul, blanco_negrita, centro, borde

    def tabla(ws, fila, titulo, cabecera, filas, col=1):
        ws.cell(fila, col, titulo).font = Font(bold=True, size=11, color=AZUL[1:])
        encabezado(ws, fila + 1, cabecera, col)
        for i, f in enumerate(filas, start=fila + 2):
            for j, v in enumerate(f):
                ws.cell(i, col + j, v).border = borde
        return fila + 2 + len(filas) + 1

    # ----- Resumen
    ws = wb.active
    ws.title = "Resumen"
    ws["A1"] = "Unidad de Asistencia Técnica · UDAT"
    ws["A1"].font = Font(bold=True, size=14, color=AZUL[1:])
    ws["A2"] = (("Informe diario" if not datos["una_hoja"] else "Hoja de ruta") + " — " + fecha_larga(datos["fecha"])
                + (f" · corte a las {timezone.localtime(datos['corte']):%H:%M}" if datos["corte"] else ""))
    ws["A3"] = f"Generado el {timezone.localtime(datos['generado_en']):%d/%m/%Y %H:%M}" + (f" por {datos['generado_por']}" if datos["generado_por"] else "")
    ws["A3"].font = Font(italic=True, color="555555")
    fila = tabla(ws, 5, "Indicadores del día", ["Hojas", "Servicios", "Con OK", "Sin OK", "Pasadas a otra brigada", "No visitadas", "En curso", "Efectividad", "Productividad", "Tiempo real (min)", "Incidencias (min)", "Tiempo disponible (min)", "Ocupación"],
                 [[t["hojas"], t["acciones"], t["ok"], t["sin_ok"], t["pasadas"], t["no_visitadas"], t["en_curso"], pct(t["efectividad_pct"]), pct(t["productividad_pct"]),
                   t["tiempo_real_min"], t["incidencias_min"], t["tiempo_disponible_min"], pct(t["ocupacion_pct"])]])
    fila = tabla(ws, fila, "Hojas de ruta", ["Hoja", "Operarios", "Servicios", "Con OK", "Sin OK", "Pasadas", "No visit.", "Efectividad", "Productividad"],
                 [[f"{h['centro']} · {h['brigada']}", ", ".join(h["operarios"]), h["acciones"], h["ok"], h["sin_ok"], h["pasadas"], h["no_visitadas"],
                   pct(h["efectividad_pct"]), pct(h["productividad_pct"])] for h in r["hojas"]])
    fila = tabla(ws, fila, "Productividad individual (la de las hojas en que participó cada operario)", ["Operario", "Hojas", "Servicios", "Con OK", "Efectividad", "Productividad"],
                 [[o["nombre"], o["hojas"], o["acciones"], o["ok"], pct(o["efectividad_pct"]), pct(o["productividad_pct"])] for o in r["por_operario"]])
    fila = tabla(ws, fila, "Por centro (según el centro de cada servicio)", ["Centro", "Servicios", "Con OK", "Efectividad", "Productividad"],
                 [[c["nombre"], c["acciones"], c["ok"], pct(c["efectividad_pct"]), pct(c["productividad_pct"])] for c in r["por_centro"]])
    fila = tabla(ws, fila, f"Pasaron a otra brigada ({r['pasadas_a_brigada']['total']})", ["Servicio", "De", "Hacia", "Estado actual"],
                 [[p["servicio"], f"{p['centro']} · {p['desde']}", p["hacia"], p["estado_actual"]] for p in r["pasadas_a_brigada"]["lista"]])
    fila = tabla(ws, fila, f"No visitadas ({r['no_visitadas']['total']})", ["Servicio", "Brigada", "Situación"],
                 [[n["servicio"], f"{n['centro']} · {n['brigada']}", "Sigue sin visitar" if n["sigue_sin_visitar"] else "Ya atendida"] for n in r["no_visitadas"]["lista"]])
    fila = tabla(ws, fila, f"Reincidencias ({r['reincidencias']['total']})", ["Servicio", "Órdenes", "Operarios que lo trabajaron"],
                 [[x["servicio"], x["ordenes"], ", ".join(x["operarios"])] for x in r["reincidencias"]["lista"]])
    tabla(ws, fila, "Incidencias por causa", ["Causa", "Tiempo (min)", "Hojas"], [[f"{i['codigo']} · {i['nombre']}", i["minutos"], i["hojas"]] for i in r["incidencias"]])
    for col, ancho in zip("ABCDEFGHIJKLM", [38, 34, 14, 14, 20, 14, 12, 14, 14, 16, 16, 20, 12]):
        ws.column_dimensions[col].width = ancho

    # ----- Hojas de ruta: mismo formato que la hoja original (bloques de 5 columnas, uno al lado del otro)
    wh = wb.create_sheet("Hojas de ruta")
    for i, h in enumerate(datos["hojas"]):
        c0 = 1 + 5 * i  # columna donde empieza el bloque
        col = lambda k: c0 + k  # noqa: E731
        wh.merge_cells(start_row=1, start_column=col(1), end_row=1, end_column=col(3))
        titulo = wh.cell(1, col(1), "HOJA DE RUTA")
        titulo.font, titulo.alignment = Font(bold=True, size=14, color=AZUL[1:]), centro
        # Centro: x-x y Operarios: x-x en lugar de las etiquetas "Centro" y "pareja" de la hoja manuscrita
        for fila_enc, texto_izq, etiqueta_der, valor_der in ((2, f"Centro: {h['centros_claves']}", "fecha", h["fecha"]),
                                                              (3, f"Operarios: {h['operarios_claves']}", "chapa", h["vehiculo"])):
            wh.merge_cells(start_row=fila_enc, start_column=col(0), end_row=fila_enc, end_column=col(2))
            wh.cell(fila_enc, col(0), texto_izq).font = Font(bold=True)
            wh.cell(fila_enc, col(3), etiqueta_der).font = Font(bold=True)
            c = wh.cell(fila_enc, col(4), valor_der)
            c.font, c.alignment = Font(bold=True), centro
            if etiqueta_der == "fecha":
                c.number_format = "DD/MM/YYYY"
            for k in range(5):
                wh.cell(fila_enc, col(k)).border = borde
        wh.merge_cells(start_row=4, start_column=col(1), end_row=4, end_column=col(2))
        for k, texto in ((0, "Trabajo"), (1, "Descripción del Trabajo"), (3, "Hora"), (4, "Cliente")):
            wh.cell(4, col(k), texto)
        for k in range(5):
            c = wh.cell(4, col(k))
            c.fill, c.font, c.alignment, c.border = azul, blanco_negrita, centro, borde
        fila = 4
        for n, a in _filas(h):
            fila += 1
            wh.merge_cells(start_row=fila, start_column=col(1), end_row=fila, end_column=col(2))
            wh.cell(fila, col(0), n).alignment = centro
            descripcion = wh.cell(fila, col(1), a["descripcion"])
            descripcion.alignment = Alignment(vertical="center", wrap_text=True)
            wh.row_dimensions[fila].height = 15 * max(1, -(-len(a["descripcion"]) // 46))  # las celdas combinadas no se ajustan solas
            wh.cell(fila, col(3), a["hora"]).alignment = centro
            wh.cell(fila, col(4), a["cliente"])
            for k in range(5):
                wh.cell(fila, col(k)).border = borde
        if fila == 4:
            wh.cell(5, col(1), "Sin acciones registradas").font = Font(italic=True, color="555555")
            fila = 5
        fila += 2
        for texto in [f"{ind}" for ind in ([f"Cerrada el {timezone.localtime(h['cerrada_en']):%d/%m/%Y %H:%M} por {h['cerrada_por']}"] if h["cerrada"] and h["cerrada_en"] else ["Hoja abierta"])]:
            wh.cell(fila, col(1), texto).font = Font(italic=True, color="555555")
            fila += 1
        for inc in h["incidencias"]:
            wh.cell(fila, col(1), f"Incidencia {inc['causa']}: {inc['inicio']} a {inc['fin']} ({minutos_a_texto(inc['minutos'])})")
            fila += 1
        for k, ancho in enumerate([8, 30, 16, 9, 18]):
            wh.column_dimensions[get_column_letter(col(k))].width = ancho
    wh.page_setup.orientation, wh.page_setup.fitToWidth, wh.page_setup.fitToHeight = "landscape", 1, 0
    wh.sheet_properties.pageSetUpPr.fitToPage = True

    # ----- Auditoría: una fila por servicio
    wa = wb.create_sheet("Auditoría")
    cab = ["Fecha", "Brigada", "Centros de la hoja", "Operarios", "Vehículo", "N.º", "Tipo de acción", "Servicio telefónico", "Tipo de mantenimiento",
           "Tipo de orden", "N.º de orden", "Trabajos (orden de trabajo)", "Centro del servicio", "Claves", "Resultado",
           "Hora de fin", "Tiempo real (min)", "Norma (min)", "Cliente (certificó OK)", "Certificada por", "Registrada por", "Hoja cerrada"]
    encabezado(wa, 1, cab)
    fila = 2
    for h in datos["hojas"]:
        for a in h["acciones"]:
            valores = [h["fecha"], h["brigada"], h["centros_claves"], h["operarios_claves"], h["vehiculo"], a["numero"], a["categoria"], a["servicio"],
                       a["tipo_mantenimiento"], a["tipo"], a["numero_orden"], ", ".join(a["trabajos"]), a["centro"], " ".join(a["claves"]), a["resultado"], a["hora"], a["tiempo_real"], a["norma"], a["cliente"], a["certificada_por"], a["creada_por"], "Sí" if h["cerrada"] else "No"]
            for j, v in enumerate(valores, start=1):
                c = wa.cell(fila, j, v)
                c.border = borde
                if j == 1:
                    c.number_format = "DD/MM/YYYY"
            fila += 1
    wa.freeze_panes = "A2"
    wa.auto_filter.ref = f"A1:{get_column_letter(len(cab))}{max(fila - 1, 1)}"
    for j, ancho in enumerate([12, 20, 12, 14, 12, 6, 16, 16, 14, 10, 12, 30, 12, 24, 22, 10, 12, 10, 24, 22, 22, 12], start=1):
        wa.column_dimensions[get_column_letter(j)].width = ancho

    for hoja_excel in (ws, wa):  # todas las hojas del libro se imprimen en horizontal, ajustadas al ancho de la página
        hoja_excel.page_setup.orientation, hoja_excel.page_setup.fitToWidth, hoja_excel.page_setup.fitToHeight = "landscape", 1, 0
        hoja_excel.sheet_properties.pageSetUpPr.fitToPage = True

    memoria = io.BytesIO()
    wb.save(memoria)
    return memoria.getvalue()


# ---------------------------------------------------------------- punto de entrada

TIPOS = {"pdf": "application/pdf", "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}


def renderizar(datos, formato="pdf"):
    """Devuelve (contenido, tipo_mime, nombre_de_archivo) a partir de los datos ya reunidos."""
    if formato not in TIPOS:
        raise ValueError("El formato debe ser pdf o xlsx.")
    contenido = _pdf(datos) if formato == "pdf" else _excel(datos)
    fecha = datos["fecha"]
    if datos["una_hoja"] and datos["hojas"]:
        h = datos["hojas"][0]
        nombre = f"UDAT_hoja_de_ruta_{fecha:%Y-%m-%d}_{_limpio(h['centros'])}_{_limpio(h['brigada'])}_{h['id']}.{formato}"
    else:
        nombre = f"UDAT_informe_diario_{fecha:%Y-%m-%d}.{formato}"
    return contenido, TIPOS[formato], nombre


def generar_informe(fecha, formato="pdf", usuario=None, brigada=None, centro=None, solo_hoja=None, ahora=None, corte=None):
    """Devuelve (contenido, tipo_mime, nombre_de_archivo). Lo usan la API y el envío automático por correo."""
    if formato not in TIPOS:
        raise ValueError("El formato debe ser pdf o xlsx.")
    return renderizar(reunir_datos(fecha, brigada, centro, solo_hoja, usuario, ahora, corte), formato)
