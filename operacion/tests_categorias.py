"""Los cuatro tipos de acción: correctivo, preventivo, orden de servicio y orden de trabajo."""
import io
from datetime import date

from django.core.exceptions import ValidationError

from catalogos.models import Clave, NormaTrabajo, TipoMantenimiento, TipoOrden, TipoTrabajo, Centro, Brigada
from .informes import generar_informe
from .models import OrdenServicio, HojaDeRuta
from .reportes import resumen_periodo
from .services import (asignar_no_visitadas, asignar_orden, asignar_pendientes, cerrar_hoja, certificar_ok, norma_accion, obtener_orden,
                       pasar_a_brigada, ordenes_pendientes)
from .tests import BaseOperacion, DIA, crear_hoja, hora
from .tests_reportes import en

CORRECTIVO, PREVENTIVO, ORDEN_SERVICIO, ORDEN_TRABAJO = "CORRECTIVO", "PREVENTIVO", "ORDEN_SERVICIO", "ORDEN_TRABAJO"


class BaseCategorias(BaseOperacion):
    def setUp(self):
        self.mppe = TipoMantenimiento.objects.get(codigo="MPPE")
        self.mpi = TipoMantenimiento.objects.get(codigo="MPI")
        self.ta = TipoOrden.objects.get(codigo="TA")
        self.poste = TipoTrabajo.objects.get(nombre="Sustitución de poste(s)")
        self.caja = TipoTrabajo.objects.get(nombre="Sustitución de caja(s) terminal(es)")
        self.h = self.hoja(DIA, operarios=[self.op1])

    def correctivo(self, numero="46 6001"):
        return obtener_orden(CORRECTIVO, self.user, numero_servicio=numero)

    def preventivo(self, numero="46 6002", tipo=None):
        return obtener_orden(PREVENTIVO, self.user, numero_servicio=numero, tipo_mantenimiento=tipo or self.mppe)

    def orden_servicio(self, numero="46 6003"):
        return obtener_orden(ORDEN_SERVICIO, self.user, numero_servicio=numero, tipo_orden=self.ta, numero_orden="7001")

    def orden_trabajo(self, numero="8001", trabajos=None):
        return obtener_orden(ORDEN_TRABAJO, self.user, numero_orden=numero, trabajos=trabajos or [self.poste])


class TiposDeAccionTests(BaseCategorias):
    def test_como_se_identifica_cada_tipo(self):
        self.assertEqual([self.correctivo().etiqueta, self.preventivo().etiqueta, self.orden_servicio().etiqueta, self.orden_trabajo().etiqueta],
                         ["46 6001", "MPPE 46 6002", "TA 7001 46 6003", "OT 8001"])

    def test_solo_se_piden_y_aceptan_los_campos_de_cada_tipo(self):
        def error(categoria, **datos):
            with self.assertRaises(ValidationError) as e:
                obtener_orden(categoria, self.user, **datos)
            return e.exception.messages[0]
        self.assertIn("el número del servicio telefónico", error(CORRECTIVO))
        self.assertIn("no corresponde indicar: el número de orden", error(CORRECTIVO, numero_servicio="46 1", numero_orden="9"))
        self.assertIn("no corresponde indicar: los trabajos", error(CORRECTIVO, numero_servicio="46 1", trabajos=[self.poste]))
        self.assertIn("el tipo de mantenimiento", error(PREVENTIVO, numero_servicio="46 1"))
        self.assertIn("no corresponde indicar: el tipo de orden", error(PREVENTIVO, numero_servicio="46 1", tipo_mantenimiento=self.mppe, tipo_orden=self.ta))
        msg = error(ORDEN_SERVICIO)
        for campo in ("el número del servicio telefónico", "el número de orden", "el tipo de orden"):
            self.assertIn(campo, msg)
        self.assertIn("no corresponde indicar: el tipo de mantenimiento", error(ORDEN_SERVICIO, numero_servicio="46 1", tipo_orden=self.ta, numero_orden="1", tipo_mantenimiento=self.mppe))
        self.assertIn("los trabajos", error(ORDEN_TRABAJO, numero_orden="1"))
        self.assertIn("no corresponde indicar: el número del servicio telefónico", error(ORDEN_TRABAJO, numero_orden="1", trabajos=[self.poste], numero_servicio="46 1"))
        self.assertIn("la categoría", error("OTRA", numero_servicio="46 1").lower())
        self.assertEqual(OrdenServicio.objects.count(), 0)  # nada se crea a medias

    def test_una_orden_abierta_se_reutiliza_y_una_restablecida_no(self):
        primera = self.correctivo()
        asignar_orden(self.h, primera, self.user)
        self.assertEqual(self.correctivo().pk, primera.pk)
        a = primera.acciones.get(); a.claves.set(Clave.objects.filter(codigo="25-R"))
        certificar_ok(a, self.user, "Cliente", ahora=hora(9))
        self.assertNotEqual(self.correctivo().pk, primera.pk)  # restablecida: se abre otra orden (reincidencia)

    def test_cada_tipo_es_una_orden_distinta_aunque_sea_el_mismo_servicio(self):
        ordenes = {self.correctivo("46 9").pk, self.preventivo("46 9").pk, self.preventivo("46 9", self.mpi).pk, self.orden_servicio("46 9").pk}
        self.assertEqual(len(ordenes), 4)
        self.assertEqual(self.preventivo("46 9").pk, self.preventivo("46 9").pk)  # el mismo preventivo sí se reutiliza

    def test_el_correctivo_exige_el_cliente_y_los_demas_tipos_no(self):
        h = self.h
        c = asignar_orden(h, self.correctivo("46 6101"), self.user); c.claves.set(Clave.objects.filter(codigo="25-R"))
        with self.assertRaises(ValidationError):
            certificar_ok(c, self.user, "", ahora=hora(9))
        p = asignar_orden(h, self.preventivo("46 6102"), self.user); p.claves.set(Clave.objects.filter(codigo="21-R"))
        certificar_ok(p, self.user, "", ahora=hora(10))  # preventivo: el cliente es opcional
        self.assertTrue(p.ok)
        o = asignar_orden(h, self.orden_servicio("46 6103"), self.user); o.claves.set(Clave.objects.filter(codigo="07-IC"))
        certificar_ok(o, self.user, "", ahora=hora(11))
        self.assertTrue(o.ok)

    def test_una_orden_de_trabajo_se_cierra_con_sus_trabajos_no_con_claves(self):
        ot = asignar_orden(self.h, self.orden_trabajo(), self.user)
        self.assertEqual(list(ot.trabajos.all()), [self.poste])  # los trabajos elegidos al asignarla
        ot.trabajos.clear()
        with self.assertRaises(ValidationError) as e:
            certificar_ok(ot, self.user, "", ahora=hora(9))
        self.assertIn("trabajo", e.exception.messages[0])
        with self.assertRaises(ValidationError):
            pasar_a_brigada(ot, self.cable, self.user, ahora=hora(9))
        ot.trabajos.set([self.poste])
        certificar_ok(ot, self.user, "", ahora=hora(10))  # sin cliente
        ot.refresh_from_db(); ot.orden.refresh_from_db()
        self.assertEqual((ot.ok, ot.orden.estado), (True, OrdenServicio.Estado.RESTABLECIDA))

    def test_la_norma_de_una_orden_de_trabajo_es_la_del_trabajo_mayor(self):
        NormaTrabajo.objects.create(tipo_trabajo=self.poste, minutos=120, vigente_desde=date(2026, 10, 1))
        NormaTrabajo.objects.create(tipo_trabajo=self.caja, minutos=90, vigente_desde=date(2026, 10, 1))
        ot = asignar_orden(self.h, self.orden_trabajo(trabajos=[self.poste, self.caja]), self.user, trabajos=[self.poste, self.caja])
        self.assertEqual(norma_accion(ot), 120)
        ot.trabajos.set([self.caja])
        self.assertEqual(norma_accion(ot), 90)
        antes = crear_hoja(fecha=date(2026, 9, 1), brigada=self.rep, centro=self.centro); antes.operarios.set([self.op2])
        previo = asignar_orden(antes, self.orden_trabajo("8002", [self.poste]), self.user, trabajos=[self.poste])
        self.assertIsNone(norma_accion(previo))  # la norma empieza el 1/10: antes no había

    def test_los_trabajos_de_una_orden_de_trabajo_no_visitada_pasan_al_dia_siguiente(self):
        asignar_orden(self.h, self.orden_trabajo(trabajos=[self.poste, self.caja]), self.user, trabajos=[self.poste, self.caja])
        cerrar_hoja(self.h, self.user, ahora=hora(15))
        manana = crear_hoja(fecha=date(2026, 10, 2), brigada=self.rep, centro=self.centro); manana.operarios.set([self.op1])
        traidas = asignar_no_visitadas(manana, self.user)
        self.assertEqual([sorted(t.nombre for t in a.trabajos.all()) for a in traidas], [sorted([self.poste.nombre, self.caja.nombre])])

    def test_pendientes_en_bloque_por_brigada_y_centro(self):
        fi = Centro.objects.get(codigo="FI")
        a = asignar_orden(self.h, self.correctivo("46 6201"), self.user); a.claves.set(Clave.objects.filter(codigo="23"))
        pasar_a_brigada(a, self.cable, self.user, ahora=hora(9))
        hoja_fi = crear_hoja(fecha=DIA, brigada=self.rep, centro=fi); hoja_fi.operarios.set([self.op2])
        b = asignar_orden(hoja_fi, self.correctivo("46 6202"), self.user); b.claves.set(Clave.objects.filter(codigo="23"))
        pasar_a_brigada(b, self.cable, self.user, ahora=hora(9))
        cable_ng = crear_hoja(fecha=date(2026, 10, 2), brigada=self.cable, centro=self.centro); cable_ng.operarios.set([self.op3])
        self.assertEqual({o.numero_servicio for o in ordenes_pendientes(self.cable, [self.centro])}, {"46 6201"})
        self.assertEqual({o.numero_servicio for o in ordenes_pendientes(self.cable)}, {"46 6201", "46 6202"})
        traidas = asignar_pendientes(cable_ng, self.user)
        self.assertEqual([t.orden.numero_servicio for t in traidas], ["46 6201"])  # solo el de su centro
        self.assertEqual(asignar_pendientes(cable_ng, self.user), [])


class ResumenesPorTipoTests(BaseCategorias):
    def escenario(self):
        from catalogos.models import Vehiculo
        NormaTrabajo.objects.create(tipo_trabajo=self.poste, minutos=120, vigente_desde=date(2026, 1, 1))
        self.h.vehiculo = Vehiculo.objects.create(chapa="B079342"); self.h.save()
        casos = [(self.correctivo("46 6001"), "25-R", hora(9)), (self.preventivo("46 6002"), "21-R", hora(10)),
                 (self.orden_servicio("46 6003"), "07-IC", hora(11)), (self.orden_trabajo("8001"), None, hora(12))]
        for orden, clave, cuando in casos:
            a = asignar_orden(self.h, orden, self.user, trabajos=[self.poste] if clave is None else None)
            if clave:
                a.claves.set(Clave.objects.filter(codigo=clave))
            certificar_ok(a, self.user, "Cliente", ahora=cuando)
        cerrar_hoja(self.h, self.user, ahora=hora(15))

    def test_desglose_por_tipo_de_accion_y_normado_de_la_orden_de_trabajo(self):
        self.escenario()
        r = resumen_periodo("dia", DIA, ahora=hora(16))
        self.assertEqual({c["codigo"]: c["acciones"] for c in r["por_categoria"]}, {"CORRECTIVO": 1, "PREVENTIVO": 1, "ORDEN_SERVICIO": 1, "ORDEN_TRABAJO": 1})
        # normado: 45 (25-R) + 35 (21-R) + 20 (07-IC) + 120 (poste) = 220 sobre 240 min reales
        self.assertEqual((r["totales"]["tiempo_normado_min"], r["totales"]["tiempo_real_min"], r["totales"]["productividad_pct"]), (220, 240, 91.7))
        ot = next(c for c in r["por_categoria"] if c["codigo"] == "ORDEN_TRABAJO")
        self.assertEqual((ot["ok"], ot["tiempo_normado_min"]), (1, 120))

    def test_por_vehiculo(self):
        self.escenario()
        r = resumen_periodo("dia", DIA, ahora=hora(16))
        self.assertEqual([(v["nombre"], v["hojas"], v["acciones"], v["productividad_pct"]) for v in r["por_vehiculo"]], [("B079342", 1, 4, 91.7)])

    def test_la_reincidencia_es_solo_de_correctivos(self):
        h2 = crear_hoja(fecha=date(2026, 10, 2), brigada=self.rep, centro=self.centro); h2.operarios.set([self.op2])
        for hoja in (self.h, h2):  # el mismo servicio con un preventivo cada día: no es reincidencia
            a = asignar_orden(hoja, obtener_orden(PREVENTIVO, self.user, numero_servicio="46 6300", tipo_mantenimiento=self.mppe), self.user)
            a.claves.set(Clave.objects.filter(codigo="21-R")); certificar_ok(a, self.user, "", ahora=hora(9) if hoja is self.h else en(2, 9))
        self.assertEqual(resumen_periodo("mes", DIA, ahora=hora(16))["reincidencias"]["total"], 0)
        for hoja in (self.h, h2):  # pero un correctivo que se vuelve a dañar sí lo es
            a = asignar_orden(hoja, obtener_orden(CORRECTIVO, self.user, numero_servicio="46 6301"), self.user)
            a.claves.set(Clave.objects.filter(codigo="25-R")); certificar_ok(a, self.user, "C", ahora=hora(10) if hoja is self.h else en(2, 10))
        self.assertEqual(resumen_periodo("mes", DIA, ahora=hora(16))["reincidencias"]["lista"][0]["servicio"], "46 6301")


class InformesPorTipoTests(BaseCategorias):
    def test_la_descripcion_de_cada_fila_lleva_los_datos_de_su_tipo(self):
        import openpyxl
        NormaTrabajo.objects.create(tipo_trabajo=self.poste, minutos=120, vigente_desde=date(2026, 1, 1))
        fi = Centro.objects.get(codigo="FI")
        h = crear_hoja(fecha=date(2026, 10, 7), brigada=self.rep, centros=[self.centro, fi]); h.operarios.set([self.op2, self.op3])
        casos = [(self.correctivo("46 6001"), ["25-R"]), (self.preventivo("46 6002"), ["21-R"]), (self.orden_servicio("46 6003"), ["07-IC"]), (self.orden_trabajo("8001"), [])]
        for n, (orden, claves) in enumerate(casos):
            a = asignar_orden(h, orden, self.user, centro=self.centro if n % 2 == 0 else fi)
            a.claves.set(Clave.objects.filter(codigo__in=claves))
        wb = openpyxl.load_workbook(io.BytesIO(generar_informe(date(2026, 10, 7), "xlsx")[0]))
        ws = wb["Hojas de ruta"]
        self.assertEqual([ws.cell(r, 2).value for r in range(5, 9)],
                         ["46 6001 25-R", "MPPE 46 6002 21-R", "TA 7001 46 6003 07-IC", "OT 8001 · Sustitución de poste(s)"])
        self.assertEqual((ws.cell(2, 1).value, ws.cell(3, 1).value), ("Centro: FI-NG", "Operarios: A2-A3"))
        aud = wb["Auditoría"]
        cab = [c.value for c in aud[1]]
        filas = [dict(zip(cab, [c.value for c in f])) for f in aud.iter_rows(min_row=2)]
        self.assertEqual([f["Tipo de acción"] for f in filas], ["Correctivo", "Preventivo", "Orden de servicio", "Orden de trabajo"])
        self.assertEqual((filas[1]["Tipo de mantenimiento"], filas[2]["Tipo de orden"], filas[2]["N.º de orden"]), ("MPPE", "TA", "7001"))
        self.assertEqual((filas[3]["Trabajos (orden de trabajo)"], filas[3]["N.º de orden"], filas[3]["Servicio telefónico"], filas[3]["Norma (min)"]),
                         ("Sustitución de poste(s)", "8001", None, 120))  # la norma sale del trabajo, aunque la acción siga abierta
